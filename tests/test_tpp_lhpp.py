"""TPP-LHPP button (tpp_lhpp.py): Offers vs. LandHoldingTransactions price check.

Serves both reports from the fake Ark site and reads back PHX_Price_Paid_Export.xlsx.
"""
from datetime import date

import pytest
from openpyxl import load_workbook

import ark_common
import tpp_lhpp

_OFFERS_URL = tpp_lhpp._REPORT_URLS[0][0]  # Offers report
_LHT_URL = tpp_lhpp._REPORT_URLS[1][0]  # LandHoldingTransactions report


@pytest.fixture
def out_dir(tmp_path, monkeypatch):
    """Redirects the hardcoded Desktop\\TPP-LHPP folder to a temp folder."""
    monkeypatch.setattr(tpp_lhpp, '_OUT_DIR', str(tmp_path))
    return tmp_path


def _export_rows(out_dir):
    """Rows of the export workbook's only sheet."""
    wb = load_workbook(out_dir / tpp_lhpp._EXPORT_FILENAME)
    return [list(r) for r in wb['Export Data'].iter_rows(values_only=True)]


def test_exports_only_offers_off_by_more_than_a_dollar(fake_ark, out_dir):
    """LHT rows are summed per offer; offers within $1 are left out; LHTs with no offer are ignored."""
    fake_ark.serve(_OFFERS_URL, [
        ['Id', 'Total Purchase Price', 'Status', 'Close Date'],
        ['OF-1', '1000', 'Closed', '2026-09-01'],   # LHT total 1000.50 -> within $1
        ['OF-2', '2000', 'Closed', '2026-09-02'],   # LHT total 1500 -> exported
        ['OF-3', '300', 'Pending', ''],             # no LHT rows -> 0 vs 300 -> exported
    ], filename=tpp_lhpp._OFFERS_FILENAME)
    fake_ark.serve(_LHT_URL, [
        ['LandHoldingTransaction: Offer Id', 'LandHoldingTransaction: Total LandHolding Value'],
        ['OF-1', '600'], ['OF-1', '400.50'],
        ['OF-2', '1000'], ['OF-2', '500'], ['OF-2', ''],
        ['OF-9', '99999'],                          # not in the Offers report -> ignored
    ], filename=tpp_lhpp._LHT_FILENAME)

    tpp_lhpp.main()

    rows = _export_rows(out_dir)
    assert rows[0] == ['Offer ID', 'Total Purchase Price', 'LHT Price Total', 'Close Date', 'Status']
    assert rows[1:] == [
        ['OF-2', 2000, 1500, '2026-09-02', 'Closed'],
        ['OF-3', 300, 0, None, 'Pending'],
    ]
    assert fake_ark.sessions_saved == 1 and fake_ark.browser_closed


def test_download_file_names_are_case_insensitive(fake_ark, out_dir):
    """Reports saved with different capitalization are still found."""
    fake_ark.serve(_OFFERS_URL, [['Id', 'Total Purchase Price', 'Status', 'Close Date'], ['OF-1', '5', '', '']],
                   filename=tpp_lhpp._OFFERS_FILENAME.upper())
    fake_ark.serve(_LHT_URL, [['LandHoldingTransaction: Offer Id', 'LandHoldingTransaction: Total LandHolding Value']],
                   filename=tpp_lhpp._LHT_FILENAME.lower())

    tpp_lhpp.main()

    assert _export_rows(out_dir)[1:] == [['OF-1', 5, 0, None, None]]


def _serve(fake_ark, offers, lht):
    """Serve an Offers report (offer rows) and an LHT report (lht rows) under the expected names."""
    fake_ark.serve(_OFFERS_URL, [['Id', 'Total Purchase Price', 'Status', 'Close Date'], *offers],
                   filename=tpp_lhpp._OFFERS_FILENAME)
    fake_ark.serve(_LHT_URL, [['LandHoldingTransaction: Offer Id', 'LandHoldingTransaction: Total LandHolding Value'],
                              *lht], filename=tpp_lhpp._LHT_FILENAME)


def test_notifies_which_exported_offers_have_a_close_date(fake_ark, out_dir, msgboxes):
    """After exporting, one message box counts and lists the exported offers that have a Close Date.
    Offers within $1 are not counted even when they have a Close Date."""
    _serve(fake_ark, [
        ['OF-1', '1000', 'Closed', '2026-09-01'],   # within $1 -> not exported, not counted
        ['OF-2', '2000', 'Closed', '2026-09-02'],   # exported, has close date
        ['OF-3', '300', 'Pending', ''],             # exported, no close date
        ['OF-4', '50', 'Closed', '2026-09-04'],     # exported, has close date
    ], [['OF-1', '1000']])

    tpp_lhpp.main()

    assert len(_export_rows(out_dir)) == 4  # header + 3 exported rows: the workbook is still written
    assert msgboxes == [('TPP-LHPP',
                         'Close Date found on 2 of 3 exported offers:\n\n'
                         'OF-2  2026-09-02\nOF-4  2026-09-04')]


def test_notifies_when_no_exported_offer_has_a_close_date(fake_ark, out_dir, msgboxes):
    """If no exported row has a Close Date, the message box says so."""
    _serve(fake_ark, [['OF-1', '1000', 'Closed', '2026-09-01'], ['OF-3', '300', 'Pending', '']],
           [['OF-1', '1000']])

    tpp_lhpp.main()

    assert msgboxes == [('TPP-LHPP', 'No Close Date on the 1 exported offer.')]


def test_notifies_when_nothing_was_exported(fake_ark, out_dir, msgboxes):
    """If every offer matches within $1 the message box says nothing was exported."""
    _serve(fake_ark, [['OF-1', '1000', 'Closed', '2026-09-01']], [['OF-1', '1000']])

    tpp_lhpp.main()

    assert _export_rows(out_dir) == [['Offer ID', 'Total Purchase Price', 'LHT Price Total', 'Close Date', 'Status']]
    assert msgboxes == [('TPP-LHPP', 'No offers were exported (every price matches within $1).')]


def test_close_date_list_is_capped(fake_ark, out_dir, msgboxes):
    """Past _MAX_LISTED offers the message box lists the first ones and counts the rest."""
    extra = 3
    count = tpp_lhpp._MAX_LISTED + extra
    _serve(fake_ark, [[f'OF-{i}', '10', 'Closed', '2026-09-01'] for i in range(count)], [])

    tpp_lhpp.main()

    [(title, text)] = msgboxes
    lines = text.split('\n')
    assert lines[0] == f'Close Date found on {count} of {count} exported offers:'
    assert lines[2:-1] == [f'OF-{i}  2026-09-01' for i in range(tpp_lhpp._MAX_LISTED)]
    assert lines[-1] == f'...and {extra} more (see {tpp_lhpp._EXPORT_FILENAME}).'


_TODAY = date(2026, 10, 8)  # pinned "today" for the Mirror two-week window


def _pin_today(monkeypatch):
    """Makes tpp_lhpp think today is _TODAY."""
    class _Date(date):
        @classmethod
        def today(cls):
            return cls(_TODAY.year, _TODAY.month, _TODAY.day)
    monkeypatch.setattr(tpp_lhpp, 'date', _Date)


def test_mirror_step_counts_mismatches_closing_in_the_next_two_weeks(fake_ark, out_dir, monkeypatch, msgboxes):
    """Only exported mismatches closing today through 14 days out are counted; the rest stay in the
    workbook without a mention. No message box is shown."""
    _pin_today(monkeypatch)
    _serve(fake_ark, [
        ['OF-TODAY', '100', 'Open', '2026-10-08'],      # today -> counted
        ['OF-SLASH', '100', 'Open', '10/12/2026'],      # other date format -> counted
        ['OF-DAY14', '100', 'Open', '2026-10-22'],      # 14 days out -> counted
        ['OF-DAY15', '100', 'Open', '2026-10-23'],      # 15 days out -> not counted
        ['OF-PAST', '100', 'Closed', '2026-10-07'],     # already closed -> not counted
        ['OF-BLANK', '100', 'Open', ''],                # no close date -> not counted
        ['OF-TBD', '100', 'Open', 'TBD'],               # unreadable date -> mentioned separately
        ['OF-MATCH', '100', 'Open', '2026-10-09'],      # within $1 -> not exported, not counted
    ], [['OF-MATCH', '100']])

    result = tpp_lhpp.mirror_step()

    assert result == (ark_common.REVIEW,
                      '3 price mismatches closing in the next two weeks, and 1 with a Close Date Mirror '
                      "couldn't read: review PHX_Price_Paid_Export.xlsx")
    assert len(_export_rows(out_dir)) == 1 + 7  # every mismatch is still in the workbook
    assert msgboxes == []


def test_mirror_step_with_no_mismatches_closing_soon(fake_ark, out_dir, monkeypatch, msgboxes):
    """Mismatches exist but none close in the next two weeks: Mirror reports no changes."""
    _pin_today(monkeypatch)
    _serve(fake_ark, [['OF-PAST', '100', 'Closed', '2026-09-01'], ['OF-BLANK', '100', 'Open', '']], [])

    assert tpp_lhpp.mirror_step() == (ark_common.NOTHING, 'No price mismatches closing in the next two weeks')
    assert msgboxes == []


def test_mirror_step_flags_unreadable_close_dates_even_with_none_closing_soon(fake_ark, out_dir, monkeypatch, msgboxes):
    """A mismatch whose Close Date can't be read might be closing soon, so it makes the step "needs
    review" instead of being dropped from the count without a word."""
    _pin_today(monkeypatch)
    _serve(fake_ark, [['OF-TBD', '100', 'Open', 'TBD'], ['OF-ODD', '100', 'Open', '10-12-26']], [])

    assert tpp_lhpp.mirror_step() == (
        ark_common.REVIEW,
        "No price mismatches closing in the next two weeks, but 2 have a Close Date Mirror couldn't read: "
        'review PHX_Price_Paid_Export.xlsx')
