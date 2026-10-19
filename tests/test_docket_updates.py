"""Docket Updates button (docket_updates.py): compares the ND and MT sheets of the Hearing Dockets
workbook with the RW cases report from Ark and writes the differences, a Cases Update file and the
Phoenix cases missing from RW.

Each test builds a small workbook with openpyxl in tmp_path ([Files] HearingDockets points at it)
and serves the RW report from the fake Ark (fake_ark fixture). The workbook reader, the rules that
pick which row of a re-listed case wins, the comparison and main()'s files and message boxes are
checked. Nothing is uploaded: ark_common.upload is replaced by a Recorder that must stay unused.
"""
from datetime import date, datetime

import openpyxl
import pytest

import ark_common
import docket_updates as du
from conftest import REPO
from fakes import Recorder, read_csv, write_csv

_TODAY = date(2026, 10, 16)  # "today" for every comparison: hearings after it are still to come
_ND_HEADER = ['Case No.', 'Case Type', 'Operator', 'Proposed Spacing Unit', 'Proposed Spacing Unit Acres',
              '# of Wells', 'Existing Spacing?', 'If yes, what is it?', 'Hearing Date', 'Date Added', 'State',
              'County', 'Project Name', 'Order Status', 'Order Number', 'Date Order Signed']  # ND / MT sheet columns
_RW_HEADER = ['Record Id', 'Applicant', 'Case Number', 'Case Status', 'Order Status', 'Case Type', 'Sections',
              'Hearing Date', 'Order Number', 'Date Order Signed']  # RW report columns after the 2026-10 update


def wb_row(number, sections='1-150N-100W-ALL', hearing=datetime(2026, 9, 24), case_type='Pooling',
           operator='Phoenix Operating LLC', status=None, order_number=None, signed=None, county='Divide',
           state='ND'):
    """One workbook row as {header: value}."""
    return {'Case No.': number, 'Case Type': case_type, 'Operator': operator, 'Proposed Spacing Unit': sections,
            'Hearing Date': hearing, 'State': state, 'County': county, 'Order Status': status,
            'Order Number': order_number, 'Date Order Signed': signed}


def make_workbook(path, nd=(), mt=(), extra_sheets=('UT', 'WY', 'Sheet2'), header=_ND_HEADER):
    """Write a Hearing Dockets workbook: ND and MT sheets from wb_row dicts, plus extra sheets that
    must be ignored (each holds a Phoenix case that would show up if it were read)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in (('ND', nd), ('MT', mt)):
        ws = wb.create_sheet(name)
        ws.append(header)
        for row in rows:
            ws.append([row.get(h) for h in header])
    for name in extra_sheets:
        ws = wb.create_sheet(name)
        ws.append(_ND_HEADER)
        ws.append([wb_row('99999').get(h) for h in _ND_HEADER])
    wb.save(path)
    return path


def rw_row(number, sections='1-150N-100W', status='', order_status='', case_type='Pooling', hearing='',
           order_number='', signed='', record_id=None, applicant='Phoenix Operating LLC'):
    """One RW report row (see _RW_HEADER)."""
    return [record_id or f'id{number}', applicant, number, status, order_status, case_type, sections, hearing,
            order_number, signed]


def compare(nd=(), rw=(), header=_RW_HEADER, mt=(), tmp_path=None, today=_TODAY):
    """Run the comparison on in-memory rows: returns (comparison rows as dicts, upload rows as dicts,
    Not in RW rows as dicts, counts)."""
    rows = []
    for sheet, wb_rows in (('ND', nd), ('MT', mt)):
        for n, r in enumerate(wb_rows, 2):
            rows.append({'sheet': sheet, 'row': n, 'number': du._text(r['Case No.']), 'type': du._text(r['Case Type']),
                         'operator': du._text(r['Operator']), 'sections': r['Proposed Spacing Unit'],
                         'hearing': du._date(r['Hearing Date']), 'state': r['State'] or sheet,
                         'county': r['County'] or '', 'status_text': du._text(r['Order Status']),
                         'status': du._status(du._text(r['Order Status'])),
                         'order_number': du._order_number(r['Order Number']), 'signed': du._date(r['Date Order Signed'])})
    rw_path = tmp_path / 'rw.csv'
    write_csv(rw_path, [header] + [r[:len(header)] for r in rw])
    records, fields = du._read_rw(str(rw_path))
    diff, uploads, missing, counts, _ = du._compare(du._cases(rows, today), records, fields)
    return ([dict(zip(du._DIFF_HEADER, r)) for r in diff], [dict(zip(['Record Id'] + fields, r)) for r in uploads],
            [dict(zip(du._MISSING_HEADER, r)) for r in missing], counts)


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    """Path of the test's Hearing Dockets workbook; [Files] HearingDockets reads it."""
    path = tmp_path / 'Hearing Dockets.xlsx'
    real = ark_common.read_setting

    def read_setting(section, key, default):
        """read_setting with [Files] HearingDockets pointed at this test's workbook."""
        return str(path) if (section, key) == ('Files', 'HearingDockets') else real(section, key, default)
    monkeypatch.setattr(ark_common, 'read_setting', read_setting)
    return path


@pytest.fixture
def no_upload(monkeypatch):
    """Recorder over ark_common.upload: Docket Updates must never call it."""
    rec = Recorder()
    monkeypatch.setattr(ark_common, 'upload', rec)
    return rec


def run(fake_ark, rw_rows, header=_RW_HEADER):
    """Serve the RW report and run main() as of _TODAY."""
    fake_ark.serve(du._REPORT_URL, [header] + [r[:len(header)] for r in rw_rows], 'RW_Case_Info.csv')
    du.main(today=_TODAY)


# ---- Reading the workbook ----

def test_reads_nd_and_mt_only(workbook, tmp_path):
    """Only the ND and MT sheets are read; UT, WY and Sheet2 are ignored. Rows without a case number
    are skipped, numeric case numbers become text and the sheet's non-breaking spaces are blanks."""
    make_workbook(workbook, nd=[wb_row(33256), wb_row(None), wb_row('33257', county='\xa0')],
                  mt=[wb_row('511-2026', state='MT')])
    rows = du._read_workbook(str(workbook))
    assert [(r['sheet'], r['number'], r['row']) for r in rows] == [('ND', '33256', 2), ('ND', '33257', 4),
                                                                   ('MT', '511-2026', 2)]
    assert rows[1]['county'] == '' and rows[0]['hearing'] == date(2026, 9, 24)


def test_missing_workbook_says_where_to_set_it(workbook, folders, fake_ark, msgboxes, no_upload):
    """No workbook at [Files] HearingDockets: an error box naming the path and the Settings row, no download."""
    du.main(today=_TODAY)
    [(title, text)] = msgboxes
    assert title == 'Docket Updates' and str(workbook) in text and 'Settings' in text
    assert fake_ark.downloads == []


def test_missing_column_is_named(workbook, folders, fake_ark, msgboxes, no_upload):
    """A sheet without a needed header stops the run with a box naming the sheet and column."""
    make_workbook(workbook, nd=[wb_row(1)], header=[h for h in _ND_HEADER if h != 'Order Status'])
    du.main(today=_TODAY)
    [(_, text)] = msgboxes
    assert 'ND sheet' in text and '"Order Status"' in text
    assert fake_ark.downloads == []


def test_missing_sheet_is_named(workbook, folders, fake_ark, msgboxes, no_upload):
    """A workbook without the MT sheet stops the run with a box saying so."""
    wb = openpyxl.Workbook()
    wb.active.title = 'ND'
    wb.active.append(_ND_HEADER)
    wb.save(workbook)
    du.main(today=_TODAY)
    assert 'no MT sheet' in msgboxes[0][1]


@pytest.mark.parametrize('text, cells, unread', [
    ('1-150N-100W-ALL', 16, []),
    ('1-150N-100W', 16, []),                       # no part means the whole section
    ('1-150N-100W-All\n\n2-150N-100W-AL', 32, []),  # mixed case, blank line, AL typo
    ('4-23N-59E-W2E2, E2W2', 8, []),               # two parts on one line: the two middle columns
    ('4-23N-59E-N2, SWSE', 9, []),
    (' 3-37N-7W - S2 ', 8, []),                     # spaces anywhere
    ('3-37N-7W-LOT1\n3-37N-7W-S2', 8, ['3-37N-7W-LOT1']),
    ('Twin Buttes', 0, ['Twin Buttes']),
])
def test_workbook_sections(text, cells, unread):
    """Proposed Spacing Unit lines become grid cells; lots and names are returned as unreadable."""
    got, bad = du._workbook_cells(text)
    assert (len(got), bad) == (cells, unread)


# ---- Which row of a case wins ----

def test_latest_hearing_gives_sub_cases_hearing_and_type(tmp_path):
    """Rows on the latest hearing date with sections are the case's groups (sub-cases) in sheet order;
    Hearing Date is the latest one; Case Type is the latest one written, "(Continued)" removed."""
    diff, uploads, _, _ = compare(tmp_path=tmp_path, nd=[
        wb_row(32170, '14-161N-97W-ALL', datetime(2026, 10, 22), '(Continued) Establish Overlapping Spacing'),
        wb_row(32170, '15-161N-97W-ALL', datetime(2026, 10, 22), None),
        wb_row(32170, '9-161N-97W-ALL', datetime(2026, 8, 27), 'ESTABLISH SPACING'),
    ], rw=[rw_row('32170-1', '14-161N-97W', 'Docketed', case_type='Establish Overlapping Spacing',
                  hearing='10/22/2026'),
           rw_row('32170-2', '15-161N-97W', 'Docketed', case_type='Establish Overlapping Spacing',
                  hearing='2026-10-22')])
    assert [(r['RW Case Number'], r['Group'], r['Rows'], r['Result']) for r in diff] == [
        ('32170-1', '1 of 2', '2', 'Match'), ('32170-2', '2 of 2', '3', 'Match')]
    assert uploads == []


def test_status_comes_from_the_latest_row_with_a_result(tmp_path):
    """A GRANTED row counts even when a later hearing re-lists the case (a note says so); its Order
    Number and Date Order Signed come with it."""
    [row], [upload], _, _ = compare(tmp_path=tmp_path, nd=[
        wb_row(5, hearing=datetime(2026, 11, 19)),
        wb_row(5, hearing=datetime(2025, 9, 26), status='GRANTED', order_number=35111, signed=datetime(2025, 10, 5)),
        wb_row(5, hearing=datetime(2025, 8, 28), status='DISMISSED'),
    ], rw=[rw_row('5', hearing='11/19/2026')])
    assert row['Differences'] == 'Case Status, Order Status, Order Number, Date Order Signed'
    assert (upload['Case Status'], upload['Order Status'], upload['Order Number'], upload['Date Order Signed']) == \
        ('Ordered', 'Granted', '35111', '10/05/2025')
    assert 'GRANTED is from the 09/26/2025 hearing; listed again for 11/19/2026' in row['Notes']


def test_filled_down_order_number_is_ignored(tmp_path, monkeypatch):
    """An Order Number on more than _FILL_DOWN_CASES cases is ignored with its Date Order Signed."""
    monkeypatch.setattr(du, '_FILL_DOWN_CASES', 2)
    nd = [wb_row(n, status='GRANTED', order_number=35278, signed=datetime(2026, 10, 5)) for n in (1, 2, 3)]
    diff, uploads, _, _ = compare(tmp_path=tmp_path, nd=nd, rw=[rw_row(str(n), status='Ordered', order_status='Granted',
                                                                       hearing='09/24/2026') for n in (1, 2, 3)])
    assert [r['Result'] for r in diff] == ['Match'] * 3 and uploads == []
    assert 'Order Number 35278 ignored (on more than 2 cases)' in diff[0]['Notes']


def test_no_result_and_a_hearing_to_come_is_docketed(tmp_path):
    """No result anywhere and the latest hearing today or later: RW should say Docketed."""
    [row], [upload], _, _ = compare(tmp_path=tmp_path, nd=[wb_row(7, hearing=datetime(2026, 10, 16))],
                                    rw=[rw_row('7', hearing='10/16/2026')])
    assert row['Differences'] == 'Case Status' and upload['Case Status'] == 'Docketed'
    assert upload['Order Status'] == ''  # Docketed says nothing about Order Status


def test_no_result_after_a_past_hearing_is_not_compared(tmp_path):
    """No recognized result and the hearing has passed: status isn't compared; the note gives the
    hearing and the status text (CONTINUED)."""
    [row], uploads, _, _ = compare(tmp_path=tmp_path, nd=[wb_row(8, hearing=datetime(2026, 10, 15), status='CONTINUED')],
                                   rw=[rw_row('8', status='Case Heard', hearing='10/15/2026')])
    assert row['Result'] == 'Match' and uploads == []
    assert 'No order status after the 10/15/2026 hearing (Order Status: CONTINUED)' in row['Notes']


@pytest.mark.parametrize('status, case_status, order_status', [
    ('GRANTED', 'Ordered', 'Granted'), ('Granted', 'Ordered', 'Granted'), ('GrANTED ', 'Ordered', 'Granted'),
    ('PARTIALLY GRANTED', 'Ordered', 'Granted'), ('Denied', 'Ordered', 'Denied'),
    ('DISMISSED', 'Dismissed', ''), ('Dimissed', 'Dismissed', ''), ('WITHDRAWN', 'Dismissed', ''),
])
def test_status_mapping(tmp_path, status, case_status, order_status):
    """Workbook results map to RW's Case Status / Order Status; a blank Order Status isn't uploaded."""
    [row], [upload], _, _ = compare(tmp_path=tmp_path, nd=[wb_row(9, status=status)],
                                    rw=[rw_row('9', status='Docketed', order_status='x', hearing='09/24/2026')])
    assert upload['Case Status'] == case_status
    assert upload['Order Status'] == (order_status or 'x')  # dismissed: RW's value is kept, never blanked


@pytest.mark.parametrize('status, hearing, case_status', [
    ('DISMISSED', datetime(2026, 9, 24), 'Dismissed'),
    ('WITHDRAWN', datetime(2026, 9, 24), 'Dismissed'),
    (None, datetime(2026, 11, 19), 'Docketed'),
])
def test_order_status_left_on_dismissed_or_docketed_is_flagged(tmp_path, status, hearing, case_status):
    """Dismissed and Docketed mean a blank Order Status: an RW value there is an Order Status
    difference with a note to clear it by hand. The upload never writes a blank, so a record whose
    only difference is that gets no upload row."""
    [row], uploads, _, _ = compare(tmp_path=tmp_path, nd=[wb_row(25, hearing=hearing, status=status)],
                                   rw=[rw_row('25', status=case_status, order_status='Granted',
                                              hearing=f'{hearing:%m/%d/%Y}')])
    assert row['Result'] == 'Differs' and row['Differences'] == 'Order Status'
    assert f'RW Order Status "Granted" should be blank ({case_status}): clear it in Ark by hand' in row['Notes']
    assert uploads == []


def test_blank_rw_order_status_on_dismissed_is_a_match(tmp_path):
    """Dismissed with RW's Order Status already blank is a Match."""
    [row], uploads, _, _ = compare(tmp_path=tmp_path, nd=[wb_row(27, status='DISMISSED')],
                                   rw=[rw_row('27', status='Dismissed', hearing='09/24/2026')])
    assert row['Result'] == 'Match' and uploads == []


@pytest.mark.parametrize('text, rw_type', [
    ('POOLING', 'Pooling'), ('(Continued) POOLING', 'Pooling'), ('Pooling (Continued)', 'Pooling'),
    ('ESTABLISH SPACING ', 'Establish Spacing'), ('TEMPORARY SPACING', 'Establish Temporary Spacing'),
    ('PERMANENT SPACING', 'Establish Permanent Spacing'), ('ADDITIONAL WELLS', 'Increased Density'),
    ('(Continued) AUTHORIZE WELLS', 'Increased Density'), ('Authorize Drilling', 'Permit to Drill'),
    ('Commingle', 'Commingling'), ('SWD', 'SWD Well Permit'), ('Permit/Protest', 'Permit Protest'),
    ('Revoke Pemit', 'Show Cause'), ('Define Field Limits', None), ('AMENDMENT-POOLING AND SPACING', None),
])
def test_type_mapping(text, rw_type):
    """Workbook Case Types map to RW types (_TYPE_MAP); the rest aren't compared."""
    assert du._rw_type(text) == rw_type


# ---- Comparing with RW ----

def test_every_field_is_compared_and_uploaded(tmp_path):
    """Each differing field is listed and uploaded with the workbook's value; unchanged fields carry
    RW's own value so the Update never blanks anything. Dates match whatever format RW writes, and
    RW's own dates are uploaded as MM/DD/YYYY like the workbook's, so the file has one date format."""
    [row], [upload], _, counts = compare(tmp_path=tmp_path, nd=[wb_row(
        10, '1-150N-100W-W2', datetime(2026, 9, 24), 'Pooling', status='DENIED', order_number='238-2026',
        signed=datetime(2026, 10, 1))], rw=[rw_row('10', '1-150N-100W', 'Docketed', '', 'Increased Density',
                                                   '2026-09-24', '237-2026', '10/1/2026', record_id='abc')])
    assert row['Result'] == 'Differs' and counts['Differs'] == 1
    assert row['Differences'] == 'Case Status, Order Status, Case Type, Sections, Order Number'
    assert row['Missing from RW'] == '' and row['Extra in RW'] == 'NE-1-150N-100W, SE-1-150N-100W'
    assert upload == {'Record Id': 'abc', 'Case Status': 'Ordered', 'Order Status': 'Denied', 'Case Type': 'Pooling',
                      'Sections': 'NW-1-150N-100W, SW-1-150N-100W', 'Hearing Date': '09/24/2026',
                      'Order Number': '238-2026', 'Date Order Signed': '10/01/2026'}


def test_unreadable_rw_date_is_noted(tmp_path):
    """An RW date in a format the script can't read differs from every workbook date; Notes names
    the value so a whole column of them is easy to spot. The upload gets the workbook's date."""
    [row], [upload], _, _ = compare(tmp_path=tmp_path, nd=[wb_row(24)], rw=[rw_row('24', hearing='Sep 24, 2026')])
    assert row['Differences'] == 'Hearing Date' and upload['Hearing Date'] == '09/24/2026'
    assert 'RW Hearing Date "Sep 24, 2026" not read as a date' in row['Notes']


def test_unreadable_rw_date_is_uploaded_as_is(tmp_path):
    """An RW date the workbook says nothing about is carried into the upload unchanged when it can't be read."""
    [upload] = compare(tmp_path=tmp_path, nd=[wb_row(26, status='GRANTED')],
                       rw=[rw_row('26', hearing='09/24/2026', signed='sometime')])[1]
    assert upload['Date Order Signed'] == 'sometime'


def test_blank_workbook_values_never_blank_rw(tmp_path):
    """A workbook row with no type, sections or order details leaves RW's values alone."""
    [row], [upload], _, _ = compare(tmp_path=tmp_path, nd=[wb_row(11, None, case_type=None, status='GRANTED')],
                                    rw=[rw_row('11', '2-150N-100W', 'Docketed', case_type='Pooling',
                                               hearing='09/24/2026', order_number='1', signed='01/02/2026')])
    assert row['Differences'] == 'Case Status, Order Status'
    assert (upload['Sections'], upload['Case Type'], upload['Order Number'], upload['Date Order Signed']) == \
        ('2-150N-100W', 'Pooling', '1', '01/02/2026')


def test_generic_establish_spacing_matches_any_spacing_type(tmp_path):
    """ESTABLISH SPACING matches RW's specific spacing types; a specific workbook type must match exactly,
    so RW's generic Establish Spacing against a specific type differs and the upload makes RW specific
    (intended, agreed 2026-10-19)."""
    diff, uploads, _, _ = compare(tmp_path=tmp_path, nd=[
        wb_row(12, case_type='ESTABLISH SPACING'), wb_row(13, case_type='Establish Permanent Spacing')],
        rw=[rw_row('12', case_type='Establish Overlapping Spacing', hearing='09/24/2026'),
            rw_row('13', case_type='Establish Spacing', hearing='09/24/2026')])
    assert [(r['Case Number'], r['Differences']) for r in diff] == [('12', ''), ('13', 'Case Type')]
    assert [u['Case Type'] for u in uploads] == ['Establish Permanent Spacing']


def test_well_types_compare_whole_sections(tmp_path):
    """For SWD, Commingling, ... the sheet gives the well's quarter-quarter and RW the section, so only
    the sections are compared; a different section still differs and uploads the whole section."""
    diff, uploads, _, _ = compare(tmp_path=tmp_path, nd=[
        wb_row(14, '23-160N-102W-SWNW', case_type='SWD'), wb_row(15, '7-162N-93W-SESW', case_type='Commingling')],
        rw=[rw_row('14', '23-160N-102W', case_type='SWD Well Permit', hearing='09/24/2026'),
            rw_row('15', '8-162N-93W', case_type='Commingling', hearing='09/24/2026')])
    assert [r['Differences'] for r in diff] == ['', 'Sections']
    assert diff[1]['Missing from RW'] == '7-162N-93W' and diff[1]['Extra in RW'] == '8-162N-93W'
    assert [u['Sections'] for u in uploads] == ['7-162N-93W']


def test_unreadable_sections_are_not_compared(tmp_path):
    """A group with an unreadable line (lots) skips the Sections check, and Notes shows the line."""
    [row], uploads, _, _ = compare(tmp_path=tmp_path, nd=[wb_row(16, '3-37N-7W-LOT1\n3-37N-7W-S2')],
                                   rw=[rw_row('16', '9-37N-7W', hearing='09/24/2026')])
    assert row['Result'] == 'Match' and uploads == []
    assert 'Unread sections: 3-37N-7W-LOT1' in row['Notes']


def test_sub_cases_pair_by_land_then_position(tmp_path):
    """Two rows on one hearing pair with RW's sub-cases: by same land first (row 3 with -1), then by position."""
    diff, _, _, _ = compare(tmp_path=tmp_path, nd=[wb_row(17, '2-150N-100W-ALL'), wb_row(17, '1-150N-100W-ALL')],
                            rw=[rw_row('17-1', '1-150N-100W', hearing='09/24/2026'),
                                rw_row('17-2', '5-150N-100W', hearing='09/24/2026')])
    assert [(r['Rows'], r['RW Case Number'], r['Differences']) for r in diff] == [
        ('2', '17-2', 'Sections'), ('3', '17-1', '')]


def test_same_land_on_one_hearing_is_one_sub_case(tmp_path):
    """Two rows of a case on one hearing with the same land are one case with two reliefs (ND 32346:
    Authorize Drilling and AUTHORIZE WELLS), not two sub-cases: one group, RW may have either type,
    and no -2 sub-case is reported Not in RW."""
    diff, uploads, missing, counts = compare(tmp_path=tmp_path, nd=[
        wb_row(22, '21-157N-92W-W2\n28-157N-92W-W2', case_type='Authorize Drilling', operator='EOG Resources, Inc.'),
        wb_row(22, '21-157N-92W-W2\n28-157N-92W-W2', case_type='AUTHORIZE WELLS', operator='EOG Resources, Inc.')],
        rw=[rw_row('22', 'W2-21-157N-92W, W2-28-157N-92W', case_type='Increased Density', hearing='09/24/2026')])
    assert [(r['Rows'], r['Group'], r['Result'], r['Differences']) for r in diff] == [('2, 3', '', 'Match', '')]
    assert uploads == [] and missing == [] and counts['Not in RW'] == 0
    assert 'Rows 2, 3 list the same land: one sub-case (Permit to Drill, Increased Density)' in diff[0]['Notes']
    assert diff[0]['Workbook Sections'] == 'NW-21-157N-92W, SW-21-157N-92W, NW-28-157N-92W, SW-28-157N-92W'


def test_same_land_rows_missing_from_rw_are_one_insert(tmp_path):
    """A Phoenix case listed twice on one hearing with the same land and missing from RW gives one
    Not in RW row, numbered as the case itself (not <number>-1 / -2)."""
    diff, _, missing, _ = compare(tmp_path=tmp_path, nd=[wb_row(23, case_type='Pooling'),
                                                         wb_row(23, case_type='ESTABLISH SPACING')], rw=[])
    assert [(r['Result'], r['Rows']) for r in diff] == [('Not in RW', '2, 3')]
    assert [m['Case Number'] for m in missing] == ['23'] and missing[0]['Sections'] == '1-150N-100W'


def test_phoenix_case_missing_from_rw(tmp_path):
    """A Phoenix case RW lacks is Not in RW (one row per sub-case, numbered <number>-k) and goes to
    Not_In_RW.csv in RW's terms; another operator's case RW lacks is ignored."""
    diff, uploads, missing, counts = compare(tmp_path=tmp_path, mt=[
        wb_row('600-2026', '6-22N-60E-ALL', datetime(2026, 11, 12), 'Permit to Drill', 'PHOENIX OPERATING', state='MT'),
        wb_row('600-2026', '7-22N-60E-ALL', datetime(2026, 11, 12), 'Permit to Drill', 'PHOENIX OPERATING', state='MT'),
        wb_row('601-2026', operator='Kraken Oil & Gas LLC', state='MT')], rw=[])
    assert [(r['Result'], r['Case Number'], r['Group']) for r in diff] == [
        ('Not in RW', '600-2026', '1 of 2'), ('Not in RW', '600-2026', '2 of 2')]
    assert missing[0] == {'Case Number': '600-2026-1', 'State': 'MT', 'Applicant': 'PHOENIX OPERATING',
                          'Case Status': 'Docketed', 'Order Status': '', 'Case Type': 'Permit to Drill',
                          'Sections': '6-22N-60E', 'Hearing Date': '11/12/2026', 'Order Number': '',
                          'Date Order Signed': ''}
    assert missing[1]['Case Number'] == '600-2026-2' and uploads == [] and counts['Not in RW'] == 2


def test_non_phoenix_case_in_rw_is_compared(tmp_path):
    """Another operator's case counts once RW has it."""
    [row], _, _, _ = compare(tmp_path=tmp_path, nd=[wb_row(18, operator='EOG Resources, Inc.')],
                             rw=[rw_row('18', hearing='09/24/2026', applicant='EOG Resources, Inc.')])
    assert row['Result'] == 'Match' and row['Workbook Operator'] == 'EOG Resources, Inc.'


def test_rw_cases_missing_from_the_workbook(tmp_path):
    """RW cases whose number isn't on the ND / MT sheets and leftover plain sub-cases are Not in
    workbook; Utah numbers and Roman-numeral field-rules sub-cases are skipped."""
    diff, _, _, counts = compare(tmp_path=tmp_path, nd=[wb_row(19, '1-150N-100W-ALL')], rw=[
        rw_row('19', hearing='09/24/2026'), rw_row('19-5', '9-150N-100W'), rw_row('19-V-1', '8-150N-100W'),
        rw_row('17422'), rw_row('2026-027-15'), rw_row('174-2024')])
    assert [(r['Result'], r['RW Case Number']) for r in diff] == [
        ('Match', '19'), ('Not in workbook', '19-5'), ('Not in workbook', '174-2024'), ('Not in workbook', '17422')]
    assert counts['Not in workbook'] == 3
    assert diff[1]['Notes'] == 'No workbook row of this case pairs with it' and diff[1]['Case Number'] == '19'
    assert diff[2]['Notes'] == 'Case not on the ND/MT sheets'


def test_rw_continuation_rows_add_sections(tmp_path):
    """A report row with a blank Record Id adds its Sections to the case above."""
    [row], _, _, _ = compare(tmp_path=tmp_path, nd=[wb_row(20, '1-150N-100W-ALL\n2-150N-100W-ALL')],
                             rw=[rw_row('20', '1-150N-100W', hearing='09/24/2026'),
                                 ['', '', '', '', '', '', '2-150N-100W', '', '', '']])
    assert row['Result'] == 'Match'


def test_report_without_new_columns_compares_the_rest(tmp_path):
    """Fields the report lacks aren't compared or uploaded."""
    header = ['Record Id', 'Applicant', 'Case Number', 'Case Status', 'Case Type', 'Sections']
    [row], [upload], _, _ = compare(tmp_path=tmp_path, header=header, nd=[wb_row(21, status='GRANTED', order_number=1)],
                                    rw=[['id21', 'P', '21', 'Docketed', 'Pooling', '1-150N-100W']])
    assert row['Differences'] == 'Case Status'
    assert list(upload) == ['Record Id', 'Case Status', 'Case Type', 'Sections']


# ---- main() ----

def test_writes_the_three_files_and_uploads_nothing(workbook, folders, fake_ark, msgboxes, no_upload):
    """A full run saves RW_Case_Info.csv, Docket_Differences.csv, Docket_Upload.csv and Not_In_RW.csv
    (UTF-8 with a BOM) and shows the counts; nothing is uploaded."""
    make_workbook(workbook, nd=[wb_row(30, status='GRANTED'), wb_row(31), wb_row(32)],
                  mt=[wb_row('700-2026', operator='Other Co', state='MT')])
    run(fake_ark, [rw_row('30', status='Docketed', hearing='09/24/2026', record_id='r30'),
                   rw_row('31', hearing='09/24/2026'), rw_row('99999')])
    out = folders['DocketUpdates']
    assert fake_ark.downloads == [du._REPORT_URL] and read_csv(out / 'RW_Case_Info.csv')[1][2] == '30'
    diff = read_csv(out / 'Docket_Differences.csv')
    assert diff[0] == du._DIFF_HEADER and [r[0] for r in diff[1:]] == ['Differs', 'Match', 'Not in RW',
                                                                         'Not in workbook']
    assert read_csv(out / 'Docket_Upload.csv') == [['Record Id'] + list(du._FIELDS),
                                                    ['r30', 'Ordered', 'Granted', 'Pooling', '1-150N-100W',
                                                     '09/24/2026', '', '']]
    assert read_csv(out / 'Not_In_RW.csv')[1][:3] == ['32', 'ND', 'Phoenix Operating LLC']
    assert (out / 'Docket_Upload.csv').read_bytes().startswith(b'\xef\xbb\xbf')
    [(title, text)] = msgboxes
    assert title == 'Docket Updates'
    assert 'Compared 3 workbook case(s) (ND/MT) with 3 RW case(s): 1 match, 1 differ' in text
    assert '1 RW record to update. Nothing was uploaded' in text and 'Not compared' not in text
    assert str(out / 'Docket_Upload.csv') in text
    assert no_upload.calls == []


def test_summary_lists_fields_the_report_lacks(workbook, folders, fake_ark, msgboxes, no_upload):
    """A report without the new columns still runs; the box names the fields not compared."""
    make_workbook(workbook, nd=[wb_row(40)])
    run(fake_ark, [['id40', 'P', '40', '', 'Pooling', '1-150N-100W']],
        header=['Record Id', 'Applicant', 'Case Number', 'Case Status', 'Case Type', 'Sections'])
    assert 'Not compared (no such column in the RW report): Order Status, Hearing Date, Order Number, ' \
           'Date Order Signed' in msgboxes[0][1]


def test_report_without_record_id_is_explained(workbook, folders, fake_ark, msgboxes, no_upload):
    """A report missing Record Id stops with a box listing the columns it has; no CSVs are written."""
    make_workbook(workbook, nd=[wb_row(41)])
    run(fake_ark, [['P', '41']], header=['Applicant', 'Case Number'])
    assert 'no Record Id column' in msgboxes[0][1] and 'Applicant, Case Number' in msgboxes[0][1]
    assert not (folders['DocketUpdates'] / 'Docket_Differences.csv').exists()


@pytest.mark.parametrize('locked', ['Docket_Differences.csv', 'Docket_Upload.csv', 'Not_In_RW.csv'])
def test_locked_csv_says_close_excel(workbook, folders, fake_ark, monkeypatch, msgboxes, no_upload, locked):
    """If Excel has one of the output CSVs open, the box names that file."""
    make_workbook(workbook, nd=[wb_row(42)])
    real = du._write_csv

    def write(path, header, rows):
        """_write_csv that fails like Excel's lock for the locked file."""
        if path.endswith(locked):
            raise PermissionError(13, 'Permission denied', path)
        real(path, header, rows)
    monkeypatch.setattr(du, '_write_csv', write)
    run(fake_ark, [])
    assert locked in msgboxes[0][1] and 'Close it in Excel' in msgboxes[0][1]


def test_defaults_match_the_settings_dialog():
    """The script's fallback folder and workbook are the Settings dialog's defaults."""
    settings = (REPO / 'settings.cpp').read_text(encoding='utf-8').replace('\\\\', '\\')
    assert f'TEXT("{du._DEFAULT_WORKBOOK}")' in settings
    assert f'TEXT("{du._DEFAULT_DIR}")' in settings
