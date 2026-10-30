"""Docket Updates Slack button (docket_updates_slack.py; bottom of column 1): exports
Cheyanne Winn's #docket-updates messages to CSV, reads the board order cases out of them and compares
them with the RW cases report from Ark.

main() runs for real against a fake Slack Web API (FakeSlack replaces the script's urlopen),
so paging, thread replies, rate-limit retries and Slack errors are exercised without a network.
When the messages hold cases, the RW report comes from the fake Ark (fake_ark fixture).
The CSVs and the finish message box are checked. The case parser and the comparison are also
tested directly on message text and small RW reports.
"""
import json
from datetime import datetime
from urllib.error import HTTPError
from urllib.parse import parse_qs, urlparse

import pytest

import docket_updates_slack as du
from openpyxl import load_workbook

from fakes import read_csv, write_csv

_OTHER_USER = 'U0OTHERUSR'  # someone else posting in the channel
_RW_HEADER = ['Record Id', 'Applicant', 'Case Number', 'Case Status', 'Case Type', 'Sections']  # columns of the RW report


def ts(year, month, day, hour, minute, second, micro='000100'):
    """Slack ts for a local send time, e.g. '1790971475.000100'."""
    return f'{int(datetime(year, month, day, hour, minute, second).timestamp())}.{micro}'


def msg(stamp, text, user=du._USER_ID, **extra):
    """One Slack message dict as conversations.history returns it."""
    return {'type': 'message', 'user': user, 'ts': stamp, 'text': text, **extra}


class _Response:
    """What urlopen returns: a context manager with read()."""

    def __init__(self, body):
        """body is the JSON-able dict Slack would send."""
        self._bytes = json.dumps(body).encode('utf-8')

    def __enter__(self):
        """Usable in `with urlopen(...) as r`."""
        return self

    def __exit__(self, *exc):
        """Nothing to close."""
        return False

    def read(self):
        """The JSON body as bytes."""
        return self._bytes


class FakeSlack:
    """Fake Slack Web API installed over docket_updates_slack.urlopen. Serves history (newest first)
    and replies (thread ts -> [parent, replies...]) in pages of the request's limit, and records
    every request as (method, params, Authorization header)."""

    def __init__(self, history=(), replies=None):
        """history: channel messages, newest first; replies: {thread ts: [parent, reply, ...]}."""
        self.history = list(history)
        self.replies = replies or {}
        self.requests = []
        self.rate_limit_next = 0  # this many upcoming requests get HTTP 429 first
        self.error = None         # when set, every call answers ok: false with this body

    def urlopen(self, request, timeout=None):
        """Answer one GET like slack.com/api would."""
        url = urlparse(request.full_url)
        method = url.path.rsplit('/', 1)[-1]
        params = {k: v[0] for k, v in parse_qs(url.query).items()}
        self.requests.append((method, params, request.get_header('Authorization')))
        if self.rate_limit_next:
            self.rate_limit_next -= 1
            raise HTTPError(request.full_url, 429, 'Too Many Requests', {'Retry-After': '7'}, None)
        if self.error:
            return _Response(dict(ok=False, **self.error))
        assert params['channel'] == du._CHANNEL_ID
        items = self.history if method == 'conversations.history' else self.replies[params['ts']]
        start = int(params.get('cursor', 0))
        end = start + int(params['limit'])
        body = {'ok': True, 'messages': items[start:end]}
        if end < len(items):
            body['response_metadata'] = {'next_cursor': str(end)}
        return _Response(body)

    def calls(self, method):
        """Params of every request made to method."""
        return [p for m, p, _ in self.requests if m == method]


@pytest.fixture
def slack(monkeypatch):
    """A FakeSlack wired into docket_updates_slack; sleeps are recorded instead of slept."""
    fake = FakeSlack()
    fake.slept = []
    monkeypatch.setattr(du, 'urlopen', fake.urlopen)
    monkeypatch.setattr(du.time, 'sleep', fake.slept.append)
    return fake


def run(folders):
    """Run main() and return the export's rows (header first)."""
    du.main()
    return read_csv(folders['DocketUpdates'] / 'Parsed_Slack_Messages.csv')


def _box(msgboxes):
    """Text of the single message box the run showed."""
    assert len(msgboxes) == 1, msgboxes
    title, text = msgboxes[0]
    assert title == 'Docket Updates Slack'
    return text


def test_exports_her_messages_newest_first_with_time_ids(slack, folders, fake_ark, msgboxes):
    """Only Cheyanne's messages are exported, newest first, each with its date/time ID and link."""
    newest = ts(2026, 10, 2, 14, 4, 35, '113259')
    oldest = ts(2026, 9, 30, 15, 18, 6, '880059')
    slack.history = [
        msg(newest, 'Board Order Updates from 10/2/2026:\nNORTH DAKOTA:\nCase 32369: GRANTED'),
        msg(ts(2026, 10, 1, 9, 0, 0), 'not hers', user=_OTHER_USER),
        msg(oldest, 'Board Order Updates from 9/30/2026:\nMONTANA:\nN/A'),
    ]
    fake_ark.serve(du._REPORT_URL, [_RW_HEADER])
    rows = run(folders)
    assert rows[0] == ['ID', 'Date', 'Time', 'Message', 'Reply To', 'Slack TS', 'Link']
    assert rows[1:] == [
        ['20261002-140435-113259', '10/02/2026', '02:04:35 PM',
         'Board Order Updates from 10/2/2026:\nNORTH DAKOTA:\nCase 32369: GRANTED', '', newest,
         f'https://phxenergy.slack.com/archives/C0881BBSFU4/p{newest.replace(".", "")}'],
        ['20260930-151806-880059', '09/30/2026', '03:18:06 PM',
         'Board Order Updates from 9/30/2026:\nMONTANA:\nN/A', '', oldest,
         f'https://phxenergy.slack.com/archives/C0881BBSFU4/p{oldest.replace(".", "")}'],
    ]
    text = _box(msgboxes)
    assert 'Saved 2 message(s)' in text and '09/30/2026 to 10/02/2026' in text


def test_sends_the_token_and_reads_only(slack, folders, msgboxes):
    """Every request carries the settings.ini token and only calls read-only methods."""
    slack.history = [msg(ts(2026, 10, 2, 8, 0, 0), 'hi')]
    run(folders)
    assert {auth for _, _, auth in slack.requests} == {'Bearer xoxp-not-a-real-token'}
    assert {m for m, _, _ in slack.requests} <= {'conversations.history', 'conversations.replies'}


def test_follows_pages_until_the_last(slack, folders, monkeypatch, msgboxes):
    """History longer than one page is read to the end through next_cursor."""
    monkeypatch.setattr(du, '_PAGE_SIZE', 2)
    slack.history = [msg(ts(2026, 9, day, 12, 0, 0), f'day {day}') for day in range(25, 20, -1)]
    rows = run(folders)
    assert [r[3] for r in rows[1:]] == ['day 25', 'day 24', 'day 23', 'day 22', 'day 21']
    assert len(slack.calls('conversations.history')) == 3


def test_includes_her_thread_replies(slack, folders, msgboxes):
    """Her replies in any thread are exported with Reply To = the parent's ID; the parent isn't duplicated."""
    parent = ts(2026, 9, 29, 14, 49, 29, '550669')
    her_reply = ts(2026, 9, 29, 16, 0, 0, '000001')
    top = msg(parent, 'Board Order Updates from 9/29/2026', reply_count=2, thread_ts=parent)
    slack.history = [top]
    slack.replies = {parent: [top,
                              msg(ts(2026, 9, 29, 15, 0, 0), 'question?', user=_OTHER_USER, thread_ts=parent),
                              msg(her_reply, 'Corrected case 32384', thread_ts=parent)]}
    rows = run(folders)
    assert [(r[3], r[4]) for r in rows[1:]] == [
        ('Corrected case 32384', '20260929-144929-550669'),
        ('Board Order Updates from 9/29/2026', ''),
    ]
    assert rows[1][6].endswith(f'?thread_ts={parent}&cid=C0881BBSFU4')


def test_replies_in_other_peoples_threads(slack, folders, msgboxes):
    """A thread started by someone else is still searched for her replies."""
    parent = ts(2026, 9, 28, 10, 0, 0)
    top = msg(parent, 'Is 32380 granted?', user=_OTHER_USER, reply_count=1, thread_ts=parent)
    slack.history = [top]
    slack.replies = {parent: [top, msg(ts(2026, 9, 28, 11, 0, 0), 'Yes', thread_ts=parent)]}
    rows = run(folders)
    assert [r[3] for r in rows[1:]] == ['Yes']


def test_broadcast_reply_is_listed_once(slack, folders, msgboxes):
    """A reply also sent to the channel shows in both history and the thread, but is exported once."""
    parent = ts(2026, 9, 28, 10, 0, 0)
    broadcast = msg(ts(2026, 9, 28, 11, 0, 0), 'Also posted to channel', subtype='thread_broadcast',
                    thread_ts=parent)
    top = msg(parent, 'Board Order Updates', reply_count=1, thread_ts=parent)
    slack.history = [broadcast, top]
    slack.replies = {parent: [top, broadcast]}
    rows = run(folders)
    assert [r[3] for r in rows[1:]] == ['Also posted to channel', 'Board Order Updates']


def test_skips_joins_and_other_system_messages(slack, folders, msgboxes):
    """Channel joins, topic changes etc. by her are not messages and are left out."""
    slack.history = [msg(ts(2026, 9, 28, 10, 0, 0), '<@U07RLNF4F44> has joined the channel',
                         subtype='channel_join'),
                     msg(ts(2026, 9, 27, 10, 0, 0), 'Board Order Updates', subtype='file_share')]
    rows = run(folders)
    assert [r[3] for r in rows[1:]] == ['Board Order Updates']


def test_slack_markup_becomes_plain_text(slack, folders, msgboxes):
    """&amp; &lt; &gt; are decoded and <url|label> links become their label."""
    slack.history = [msg(ts(2026, 9, 28, 10, 0, 0),
                         'PETRO-HUNT &amp; SONS &lt;N/A&gt; see <https://x.example/32499|order 32499> '
                         'or <https://x.example/a>')]
    rows = run(folders)
    assert rows[1][3] == 'PETRO-HUNT & SONS <N/A> see order 32499 or https://x.example/a'


def test_ids_are_unique_within_the_same_second(slack, folders, msgboxes):
    """Two messages sent in the same second get different IDs (Slack's microseconds)."""
    slack.history = [msg(ts(2026, 9, 28, 10, 0, 0, '000002'), 'b'), msg(ts(2026, 9, 28, 10, 0, 0, '000001'), 'a')]
    rows = run(folders)
    assert [r[0] for r in rows[1:]] == ['20260928-100000-000002', '20260928-100000-000001']


def test_waits_and_retries_when_rate_limited(slack, folders, msgboxes):
    """HTTP 429 waits Retry-After seconds and retries instead of failing."""
    slack.history = [msg(ts(2026, 9, 28, 10, 0, 0), 'hi')]
    slack.rate_limit_next = 2
    rows = run(folders)
    assert [r[3] for r in rows[1:]] == ['hi']
    assert slack.slept == [7, 7]


def test_gives_up_after_too_many_rate_limits(slack, folders, msgboxes):
    """Endless 429s end in an error (the __main__ handler then shows the failure box)."""
    slack.rate_limit_next = du._MAX_RETRIES + 1
    with pytest.raises(HTTPError):
        du.main()


def test_slack_error_is_shown(slack, folders, msgboxes):
    """ok: false from Slack (e.g. a token without channels:history) explains itself and writes nothing."""
    slack.error = {'error': 'missing_scope', 'needed': 'channels:history'}
    du.main()
    text = _box(msgboxes)
    assert 'missing_scope' in text and 'channels:history' in text
    assert not (folders['DocketUpdates'] / 'Parsed_Slack_Messages.csv').exists()
    assert not (folders['DocketUpdates'] / 'Slack_Docket_Differences.xlsx').exists()


@pytest.mark.parametrize('token', [None, '   '])
def test_missing_token_explains_setup(slack, folders, monkeypatch, msgboxes, token):
    """No [SlackCredentials] token (or a blank one) shows the setup steps and calls nothing."""
    def read_slack_token():
        """Stand-in for a settings.ini with no / a blank token."""
        if token is None:
            raise KeyError('SlackCredentials')
        return token
    monkeypatch.setattr(du.ark_common, 'read_slack_token', read_slack_token)
    du.main()
    text = _box(msgboxes)
    assert '[SlackCredentials]' in text and 'channels:history' in text
    assert slack.requests == []


def test_locked_csv_says_close_excel(slack, folders, monkeypatch, msgboxes):
    """If Excel has the CSV open, the box says to close it instead of failing silently."""
    slack.history = [msg(ts(2026, 9, 28, 10, 0, 0), 'hi')]

    def locked(path, header, rows):
        """Stand-in for _write_csv on a file Excel has open."""
        raise PermissionError(13, 'Permission denied', path)
    monkeypatch.setattr(du, '_write_csv', locked)
    du.main()
    assert 'Close it in Excel' in _box(msgboxes)


def test_no_messages_still_writes_the_header(slack, folders, msgboxes):
    """An empty channel gives a header-only CSV and says so."""
    rows = run(folders)
    assert rows == [du._OUTPUT_HEADER]
    assert 'No messages' in _box(msgboxes)


def test_csv_opens_in_excel_with_em_dashes(slack, folders, fake_ark, msgboxes):
    """The CSV starts with a UTF-8 BOM so Excel reads em dashes correctly."""
    slack.history = [msg(ts(2026, 9, 28, 10, 0, 0), 'Case 33150 — GRANTED')]
    fake_ark.serve(du._REPORT_URL, [_RW_HEADER])
    run(folders)
    raw = (folders['DocketUpdates'] / 'Parsed_Slack_Messages.csv').read_bytes()
    assert raw.startswith(b'\xef\xbb\xbf')
    assert 'Case 33150 — GRANTED' in raw.decode('utf-8-sig')


def test_default_folder_is_desktop_docket_updates(slack, monkeypatch, tmp_path, msgboxes):
    """With no [Folders] DocketUpdates, the CSV goes to Desktop\\Docket Updates, created if missing
    (the run itself is redirected to tmp_path)."""
    assert du._DEFAULT_DIR == r'C:\Users\Ethan Mesecher\Desktop\Docket Updates'
    monkeypatch.setattr(du, '_DEFAULT_DIR', str(tmp_path / 'Docket Updates'))
    monkeypatch.setattr(du.ark_common, 'read_folder', lambda key, default: default)
    du.main()
    assert read_csv(tmp_path / 'Docket Updates' / 'Parsed_Slack_Messages.csv') == [du._OUTPUT_HEADER]


# ---- Reading cases out of the messages ----

_EXAMPLE = '''*Board Order Updates from 9/30/2026:*
*NORTH DAKOTA:*
 *Case 32275 (CONTINENTAL RESOURCES, INC): Pooling: GRANTED*
*19-155N-97W-ALL*
*30-155N-97W-ALL*
*31-155N-97W-ALL*
*24-155N-98W-ALL*

*MONTANA:*
*N/A*'''  # a board order message in the usual format


def export_row(text, day=30, micro='000100', link='https://slack.example/p1'):
    """One Parsed_Slack_Messages.csv row (see _OUTPUT_HEADER) holding message text, posted 9/<day>/2026."""
    return [f'202609{day:02d}-100000-{micro}', f'09/{day:02d}/2026', '10:00:00 AM', text, '', '1.1', link]


def test_reads_case_fields_from_a_message():
    """Number, applicant, type and status come out of the usual format; the section lines are skipped."""
    [case] = du._parse_message(_EXAMPLE)
    assert (case['number'], case['numbers'], case['applicant'], case['type'], case['status']) == \
        ('32275', ['32275'], 'CONTINENTAL RESOURCES, INC', 'Pooling', 'GRANTED')
    assert case['notes'] == []


@pytest.mark.parametrize('line, number, applicant, case_type, status', [
    ('Case 31935(PHOENIX OPERATING LLC): Pooling and Spacing AMENDMENT: DISMISSED',
     '31935', 'PHOENIX OPERATING LLC', 'Pooling and Spacing AMENDMENT', 'DISMISSED'),
    ('Case 32078 (CONTINENTAL RESOURCES, INC.) Pooling and Spacing: DISMISSED',  # no colon after the applicant
     '32078', 'CONTINENTAL RESOURCES, INC.', 'Pooling and Spacing', 'DISMISSED'),
    ('Case 32843 (ENERPLUS RESOURCES (USA) CORP): Pooling and Spacing AMENDMENT: GRANTED',  # nested parentheses
     '32843', 'ENERPLUS RESOURCES (USA) CORP', 'Pooling and Spacing AMENDMENT', 'GRANTED'),
    ('Cases 32952 ( FORMENTERA OPERATIONS LLC): Pooling and Spacing: GRANTED',
     '32952', 'FORMENTERA OPERATIONS LLC', 'Pooling and Spacing', 'GRANTED'),
    ('* *Case 32591( KRAKEN OPERATING LLC) * *Pooling and Spacing: GRANTED*',  # stray asterisks
     '32591', 'KRAKEN OPERATING LLC', 'Pooling and Spacing', 'GRANTED'),
    ('Case 32984 (FORMENTERA OPERATIONS, LLC) : Pooling and Spacing : GRANTED',
     '32984', 'FORMENTERA OPERATIONS, LLC', 'Pooling and Spacing', 'GRANTED'),
    ('Case 31915 (OASIS): Pooling and Spacing: ~GRANTED~ DISMISSED',  # struck out in Slack
     '31915', 'OASIS', 'Pooling and Spacing', 'DISMISSED'),
    ('Case 31414 (WGO Resources, LLC): Pooling and Spacing: GRANTED141', '31414', 'WGO Resources, LLC',
     'Pooling and Spacing', 'GRANTED'),
    ('Case 31777 (Devon): Pooling and Spacing: Granted', '31777', 'Devon', 'Pooling and Spacing', 'GRANTED'),
    ('Case 33001 (EOG): Pooling: GARNTED', '33001', 'EOG', 'Pooling', 'GRANTED'),  # typo seen in Slack
    ('Case 32373 (HESS BAKKEN INVESTMENTS II, LLC): Pooling and Spacing AMENDMENT: GRANTED:',
     '32373', 'HESS BAKKEN INVESTMENTS II, LLC', 'Pooling and Spacing AMENDMENT', 'GRANTED'),
    ('Case 32279 (PHOENIX OPERATING LLC): Pooling and Spacing AMENDMENT:',  # no status yet
     '32279', 'PHOENIX OPERATING LLC', 'Pooling and Spacing AMENDMENT', ''),
    ('Case 283-2025 (KRAKEN) : AMENDMENT – CONTESTED: PARTIALLY GRANTED – 1YR EXT. INSTEAD OF 2YR',
     '283-2025', 'KRAKEN', 'AMENDMENT – CONTESTED', 'PARTIALLY GRANTED'),
    ('Case 33114 (SOGC, INC.(?)): Pooling and Spacing AMENDMENT: DISMISSED',
     '33114', 'SOGC, INC.(?)', 'Pooling and Spacing AMENDMENT', 'DISMISSED'),
])
def test_header_variants(line, number, applicant, case_type, status):
    """The header formats seen in #docket-updates all give the right fields."""
    header = du._parse_header(' '.join(line.replace('*', ' ').split()))
    assert (header['number'], header['applicant'], header['type'], header['status']) == \
        (number, applicant, case_type, status)


def test_header_with_two_numbers_and_an_extra_note():
    """'196-2025 / 32-2025 FED' looks up both numbers; an extra (note) before the applicant is kept as a note."""
    header = du._parse_header('Case 196-2025 / 32-2025 FED (CONTINENTAL) : Pooling and Spacing: GRANTED')
    assert header['number'] == '196-2025 / 32-2025 FED' and header['numbers'] == ['196-2025', '32-2025']
    header = du._parse_header('Case 31814 (29245?) (WHITING): Pooling & Spacing: GRANTED')
    assert (header['applicant'], header['notes']) == ('WHITING', ['(29245?)'])


def test_odd_status_text_is_kept_as_a_note():
    """Extra words around the status ('PARTIALLY GRANTED – 1YR EXT...') go to the notes."""
    header = du._parse_header('Case 283-2025 (KRAKEN) : AMENDMENT: PARTIALLY GRANTED – 1YR EXT.')
    assert header['notes'] == ['Status text: PARTIALLY GRANTED – 1YR EXT.']


def test_lines_that_are_not_cases_are_skipped():
    """Board headings, N/A, No New Updates and chit-chat give no cases."""
    text = ('Board Order Updates from 9/24/2026:\nNORTH DAKOTA:\nNo New Updates\nMONTANA:\nN/A\n'
            "Sorry, y'all, forgot to post these yesterday\nCorrected case 32384")
    assert du._parse_message(text) == []


def test_status_on_the_line_below_the_header():
    """A header ending in ':' with GRANTED on the next line takes that status."""
    [case] = du._parse_message('Case 32601 (OASIS): Authorize Wells AMENDMENT:\nGRANTED\n(3 WELLS)\n1-150N-100W-ALL')
    assert case['status'] == 'GRANTED' and case['notes'] == []
    [case] = du._parse_message('*Case 32601 (OASIS): Authorize Wells AMENDMENT:*\n*\n*GRANTED*')  # blank line between
    assert case['status'] == 'GRANTED'


def test_status_word_further_down_is_ignored():
    """A status word in a remark below the sections is not taken as the case's status."""
    [case] = du._parse_message('Case 32279 (PHOENIX OPERATING LLC): Pooling and Spacing AMENDMENT:\n'
                               '1-150N-100W-ALL\nPrior application DENIED')
    assert case['status'] == '' and case['notes'] == []


def test_sections_and_states_are_not_read():
    """State headings, section lines, separators, unit names and remarks inside a case are skipped:
    nothing about land or state is kept, only the header fields and the well count."""
    text = ('*NORTH DAKOTA:*\n*Case 32939 (CONTINENTAL): Pooling and Spacing AMENDMENT : GRANTED*\n'
            '*(4 Wells-FEDOR UNIT)*\n*1-150N-100W-ALL*\n*2-150N-100W-ALL*\n*-*\n*(5 Wells)-(CHEETAH UNIT)*\n'
            '*3-150N-100W-ALL*\n*WILLOW GRAY*\n*MONTANA:*\n13-28N-57E-ALL\n'
            '*Case 281-2025 (COYOTE RESOURCES) : Pooling & Spacing: GRANTED*\n'
            ' 13-28N-57E-ALL (NW1/4 Contains Indian Trust Minerals)\n5-22N-59-S2')
    nd, mt = du._parse_message(text)
    assert (nd['number'], nd['status'], mt['number'], mt['status']) == ('32939', 'GRANTED', '281-2025', 'GRANTED')
    assert nd['notes'] == mt['notes'] == []
    assert not {'state', 'groups'} & (nd.keys() | mt.keys())


@pytest.mark.parametrize('lines, wells', [
    ('(4 WELLS)\n13-144N-95W-ALL\n(4 WELLS)\n24-144N-95W-ALL', 8),  # several groups are added up
    ('*(4 Wells-FEDOR UNIT)*\n*(5 Wells)-(CHEETAH UNIT)*', 9),     # unit names after the count
    ('(1 Well)\n1-150N-100W-ALL', 1),                               # singular
    ('2 wells\n1-150N-100W-ALL', 2),                                # no parentheses
    ('1-150N-100W-ALL\nWILLOW GRAY', None),                         # no (N Wells) line: blank
    ('1-150N-100W-ALL\nPrior application for 3 wells DENIED', None),  # a count inside a remark isn't one
])
def test_well_count_sums_the_wells_lines(lines, wells):
    """A case's well count is the sum of its (N Wells) lines, None when it has none."""
    [case] = du._parse_message(f'Case 32840 (ENERPLUS): Pooling: GRANTED\n{lines}')
    assert case['wells'] == wells


def test_wells_lines_belong_to_the_case_above():
    """Each case counts only the (N Wells) lines below its own header; one above every case is ignored,
    and a wells line right below a header with no status isn't taken as the status."""
    first, second = du._parse_message('(7 Wells)\nCase 1 (X): Pooling:\n(2 Wells)\nCase 2 (Y): Pooling: GRANTED\n'
                                      '(3 Wells)\n(1 Well)')
    assert (first['wells'], first['status'], second['wells']) == (2, '', 4)


def test_newest_post_of_a_case_wins():
    """A case posted again later is taken from the newest message; the older post's status is noted."""
    new = export_row('Case 31954 (HESS): Pooling: DISMISSED\n1-150N-100W-ALL', day=17, link='new')
    old = export_row('Case 31954 (HESS): Pooling: GRANTED\n1-150N-100W-ALL', day=9, link='old')
    [case] = du._cases_from([new, old])
    assert (case['status'], case['link'], case['posted']) == ('DISMISSED', 'new', '09/17/2026')
    assert case['notes'] == ['Also posted 09/09/2026: GRANTED']


def test_newer_post_with_only_sections_does_not_win():
    """A re-post with just the number and section lines says nothing that's compared, so the older post is kept."""
    new = export_row('Case 31954:\n1-150N-100W-ALL', day=20, link='new')
    old = export_row('Case 31954 (HESS): Pooling: GRANTED', day=9, link='old')
    [case] = du._cases_from([new, old])
    assert (case['status'], case['link']) == ('GRANTED', 'old')
    assert case['notes'] == ['Also posted 09/20/2026: no status']


def test_blanks_are_filled_from_the_newest_post_that_has_them():
    """Each blank field of the winning post comes from the newest older post that has it; filled fields stay."""
    new = export_row('Case 31954 (HESS):', day=20, link='new')
    mid = export_row('Case 31954 (HESS CO): Pooling: DENIED', day=17)
    old = export_row('Case 31954 (HESS CORP): Pooling and Spacing: GRANTED', day=9)
    [case] = du._cases_from([new, mid, old])
    assert (case['status'], case['type'], case['applicant'], case['link']) == ('DENIED', 'Pooling', 'HESS', 'new')
    assert case['notes'] == ['Also posted 09/17/2026: DENIED', 'Also posted 09/09/2026: GRANTED',
                             'Status from the 09/17/2026 post', 'Type from the 09/17/2026 post']


def test_unrecognized_status_is_not_replaced():
    """A status that is written but not a board result (CONTINUED) is info, not a blank: it isn't filled."""
    new = export_row('Case 31954 (HESS): Pooling: CONTINUED', day=17)
    old = export_row('Case 31954 (HESS): Pooling: GRANTED', day=9)
    [case] = du._cases_from([new, old])
    assert case['status'] == ''
    assert case['notes'] == ['Status text: CONTINUED', 'Also posted 09/09/2026: GRANTED']


def test_case_listed_twice_in_one_message_is_one_case():
    """Two blocks of the same case in one message give one case, noted; the second fills the first's
    blanks and both blocks' wells are counted."""
    text = 'Case 32742 (X): Pooling:\n(2 Wells)\nCase 32742 (X): Pooling: GRANTED\n(3 Wells)'
    [case] = du._cases_from([export_row(text)])
    assert (case['status'], case['wells']) == ('GRANTED', 5)
    assert case['notes'] == ['Listed more than once in this message']
    [case] = du._cases_from([export_row('Case 32742 (X): Pooling: GRANTED\nCase 32742 (X): Pooling:\n(3 Wells)')])
    assert case['wells'] == 3  # only one block has a count


def test_well_count_is_filled_from_an_older_post():
    """A newer post without (N Wells) lines takes the well count of the newest older post that has one;
    a newer post with only a well count still wins over an older one."""
    new = export_row('Case 31954 (HESS): Pooling: DISMISSED', day=17, link='new')
    old = export_row('Case 31954 (HESS): Pooling: GRANTED\n(3 Wells)', day=9)
    [case] = du._cases_from([new, old])
    assert (case['link'], case['wells']) == ('new', 3)
    assert case['notes'] == ['Also posted 09/09/2026: GRANTED', 'Well Count from the 09/09/2026 post']
    new = export_row('Case 31954:\n(4 Wells)', day=20, link='new')
    [case] = du._cases_from([new, old])
    assert (case['link'], case['wells'], case['status']) == ('new', 4, 'GRANTED')


def test_later_mention_does_not_replace_the_post():
    """A later reply that only mentions the case ("Case 32384 was corrected...") has no status, type or
    applicant, so the real post is kept and the mention is noted."""
    mention = export_row('Case 32384 was corrected, see above', day=17, link='reply')
    post = export_row('Case 32384 (HESS): Pooling: GRANTED', day=9, link='post')
    [case] = du._cases_from([mention, post])
    assert (case['status'], case['link']) == ('GRANTED', 'post')
    assert case['notes'] == ['Also posted 09/17/2026 as 32384 was corrected, see above: no status']


def test_posts_sharing_a_number_are_one_case():
    """'196-2025' posted after '196-2025 / 32-2025 FED' is the same case, and both numbers are still looked up."""
    new = export_row('Case 196-2025 (CONTINENTAL): Pooling and Spacing: DISMISSED', day=17)
    old = export_row('Case 196-2025 / 32-2025 FED (CONTINENTAL): Pooling and Spacing: GRANTED', day=9)
    [case] = du._cases_from([new, old])
    assert (case['number'], case['numbers'], case['status']) == ('196-2025', ['196-2025', '32-2025'], 'DISMISSED')
    assert case['notes'] == ['Also posted 09/09/2026 as 196-2025 / 32-2025 FED: GRANTED']


# ---- Comparing with RW ----

def rw_row(number, status='', case_type='Pooling', record_id=None, applicant='Phoenix Operating LLC',
           sections='1-150N-100W'):
    """One RW report row (see _RW_HEADER). Sections is in the real report but never compared."""
    return [record_id or f'REC-{number}', applicant, number, status, case_type, sections]


def compare(tmp_path, texts, rw_rows, header=_RW_HEADER):
    """Compare messages (newest first) with an RW report; returns (rows as dicts with Result, Result counts)."""
    path = write_csv(str(tmp_path / 'rw.csv'), [header] + rw_rows)
    rw, optional = du._read_rw(path)
    cases = du._cases_from([export_row(t, micro=f'{i:06d}') for i, t in enumerate(texts)])
    return du._compare(cases, rw, optional)


def not_in_rw(tmp_path, texts, rw_rows):
    """Slack_Not_In_RW.csv rows (as dicts) for messages (newest first) compared with an RW report."""
    rows, _ = compare(tmp_path, texts, rw_rows)
    return [dict(zip(du._MISSING_HEADER, r)) for r in du._missing_rows(rows)]


def read_differences(path):
    """(rows, highlighted, sheet) of the differences workbook: cell values row by row, header first
    (blanks as ''), and the (row index, column name) of every yellow cell."""
    ws = load_workbook(path).active
    cells = list(ws.iter_rows())
    rows = [['' if c.value is None else c.value for c in r] for r in cells]
    highlighted = {(i, rows[0][j]) for i, r in enumerate(cells) for j, c in enumerate(r)
                   if c.fill.fill_type == 'solid' and c.fill.fgColor.rgb.endswith('FFFF00')}
    return rows, highlighted, ws


_POOLING = 'Case 33068 (PHOENIX OPERATING LLC): Pooling: {status}\n4-158N-96W-ALL\n9-158N-96W-ALL'  # case with section lines


def test_matching_case_is_a_match(tmp_path):
    """Same status (GRANTED = Ordered): Match, with both sides filled in."""
    [row], counts = compare(tmp_path, [_POOLING.format(status='GRANTED')], [rw_row('33068', 'Ordered')])
    assert (row['Result'], row['Differences'], row['Record Id'], row['RW Case Number']) == \
        ('Match', '', 'REC-33068', '33068')
    assert (row['Slack Status'], row['RW Case Status']) == ('GRANTED', 'Ordered')
    assert counts == {'Match': 1}


def test_type_and_applicant_are_not_compared_or_listed(tmp_path):
    """A different Case Type or applicant doesn't make a record differ (no note either), and the
    differences workbook has no applicant or type columns."""
    [row], _ = compare(tmp_path, ['Case 1 (EOG): Pooling: GRANTED'],
                       [rw_row('1', 'Ordered', 'Commingling', applicant='Hess Corp')])
    assert (row['Result'], row['Notes']) == ('Match', '')
    assert not [h for h in du._DIFF_HEADER if 'Applicant' in h or 'Type' in h]


def test_sections_and_state_are_not_compared_or_listed(tmp_path):
    """RW sections that differ from the message's don't make a difference, and the file has no
    section, state or well group columns."""
    [row], _ = compare(tmp_path, ['NORTH DAKOTA:\n' + _POOLING.format(status='GRANTED')],
                       [rw_row('33068', 'Ordered', sections='16-158N-96W')])
    assert row['Result'] == 'Match' and row['Notes'] == ''
    assert not [h for h in du._DIFF_HEADER if 'Section' in h or h in ('State', 'Group', 'Missing from RW')]


def test_well_count_is_not_compared_without_the_rw_column(tmp_path):
    """Without a Number of Wells column the Slack Well Count is only listed (blank when Slack gives none)."""
    rows, _ = compare(tmp_path, [_GROUPED, _POOLING.format(status='GRANTED')],
                      [rw_row('32840-1', 'Ordered', 'Establish Spacing'), rw_row('32840-2', 'Ordered', 'Establish Spacing'),
                       rw_row('33068', 'Ordered')])
    assert [(r['RW Case Number'], r['Slack Well Count'], r['RW Well Count'], r['Result']) for r in rows] == \
        [('32840-1', 8, '', 'Match'), ('32840-2', 8, '', 'Match'), ('33068', '', '', 'Match')]


_WELLS_HEADER = _RW_HEADER + ['Number of Wells']  # RW report with the well count column


@pytest.mark.parametrize('slack_wells, rw_wells, differs', [
    ('\n(3 Wells)', '3', False),
    ('\n(3 Wells)', '3.0', False),   # Ark may export a number with a decimal
    ('\n(3 Wells)', '4', True),
    ('\n(3 Wells)', '', True),       # RW not filled in yet
    ('\n(3 Wells)', 'three', True),  # not a number
    ('', '5', False),                # no (N Wells) line in Slack: nothing to compare
])
def test_well_count_is_compared_with_rw(tmp_path, slack_wells, rw_wells, differs):
    """Slack's well count must equal RW's Number of Wells; RW Well Count shows RW's value (a number when readable)."""
    [row], _ = compare(tmp_path, [_POOLING.format(status='GRANTED') + slack_wells],
                       [rw_row('33068', 'Ordered') + [rw_wells]], header=_WELLS_HEADER)
    assert row['Differences'] == ('Well Count' if differs else '')
    assert row['RW Well Count'] == (int(float(rw_wells)) if rw_wells[:1].isdigit() else rw_wells)


def test_well_column_is_found_ignoring_case(tmp_path):
    """A report column named 'number of wells' (any case, stray spaces) is the well count column."""
    [row], _ = compare(tmp_path, [_POOLING.format(status='GRANTED') + '\n(2 Wells)'],
                       [rw_row('33068', 'Ordered') + ['5']], header=_RW_HEADER + [' number of WELLS '])
    assert (row['Differences'], row['RW Well Count']) == ('Well Count', 5)


def test_well_count_of_several_records_is_their_total(tmp_path):
    """Slack's count is for the whole case, so with several RW sub-cases it is compared with their total;
    when that differs every record of the case differs and a note gives the total (plus the sub-case
    note: such a case is left out of the upload)."""
    rows, _ = compare(tmp_path, [_GROUPED], [rw_row('32840-1', 'Ordered', 'Establish Spacing') + ['4'],
                                             rw_row('32840-2', 'Ordered', 'Establish Spacing') + ['4']],
                      header=_WELLS_HEADER)
    assert [r['Result'] for r in rows] == ['Match', 'Match']
    rows, _ = compare(tmp_path, [_GROUPED], [rw_row('32840-1', 'Ordered', 'Establish Spacing') + ['4'],
                                             rw_row('32840-2', 'Ordered', 'Establish Spacing') + ['']],
                      header=_WELLS_HEADER)
    assert [(r['Differences'], r['Notes']) for r in rows] == \
        [('Well Count', "Slack's 8 wells compared with the 2 RW records' total of 4; "
                        'Not in Slack_Docket_Upload.csv (sub-cases): review by hand')] * 2


@pytest.mark.parametrize('slack, rw, differs', [
    ('GRANTED', 'Ordered', False),
    ('GRANTED', '', True),            # RW not updated yet
    ('GRANTED', 'Docketed', True),
    ('PARTIALLY GRANTED', 'Ordered', False),
    ('DENIED', 'Ordered', False),
    ('DISMISSED', 'Dismissed', False),
    ('DISMISSED', 'Ordered', True),
    ('', 'Docketed', False),          # no status in Slack: nothing to compare
])
def test_status_mapping(tmp_path, slack, rw, differs):
    """Slack results map to RW Case Status: GRANTED / PARTIALLY GRANTED / DENIED = Ordered, DISMISSED = Dismissed."""
    [row], _ = compare(tmp_path, [_POOLING.format(status=slack)], [rw_row('33068', rw)])
    assert row['Differences'] == ('Status' if differs else '')


def test_order_status_is_compared_when_the_report_has_it(tmp_path):
    """With an Order Status column, GRANTED needs Granted and DISMISSED needs it blank."""
    header = _RW_HEADER + ['Order Status']
    rows, _ = compare(tmp_path, [_POOLING.format(status='GRANTED'), 'Case 33069 (X): Pooling: DISMISSED'],
                      [rw_row('33068', 'Ordered') + ['Denied'], rw_row('33069', 'Dismissed') + ['']], header=header)
    assert [(r['Case Number'], r['Differences'], r['RW Order Status']) for r in rows] == \
        [('33068', 'Order Status', 'Denied'), ('33069', '', '')]


@pytest.mark.parametrize('slack_type, rw_type, differs', [
    ('Pooling', 'Pooling', False),
    ('Pooling', 'Establish Spacing', True),
    ('Pooling and Spacing', 'Pooling', False),
    ('Pooling & Spacing', 'Establish Temporary Spacing', False),
    ('Pooling and Spacing AMENDMENT', 'Establish Overlapping Spacing', False),
    ('Pooling and Spacing AMENDMENT', 'Increased Density', False),
    ('Pooling and Spacing AMENDMENT', 'Commingling', True),
    ('Pooling and Spacing AMENDMENT', 'Pooling', True),
    ('Spacing and Pooling', 'Establish Spacing', False),          # either word order
    ('Spacing & Pooling AMENDMENT', 'Increased Density', False),
    ('Pooling/Spacing', 'Establish Spacing', False),
    ('Spacing and AMENDMENTS', 'Establish Permanent Spacing', False),
    ('Authorize Wells AMENDMENT', 'Increased Density', False),
    ('Increased Capacity', 'Increased Density', False),
    ('Capacity Increase', 'Increased Density', False),
    ('Authorize Drilling EXCEPTION', 'Permit to Drill', False),
    ('Revoke Permits', 'Show Cause', False),
    ('Suspend and Revoke', 'Pooling', True),
])
def test_type_mapping(slack_type, rw_type, differs):
    """Slack case types map to the RW Case Types listed in _TYPE_RULES (no longer compared with RW,
    but still how a header tells a type from a status)."""
    assert (rw_type.lower() not in {t.lower() for t in du._rw_types(slack_type) or ()}) == differs


def test_unmapped_type_has_no_rw_type():
    """Redefine Field Limits has no RW type."""
    assert du._rw_types('Redefine Field Limits-GRASSY BUTTE') is None


def test_rw_case_notes_are_attached_not_compared(tmp_path):
    """RW's Case Notes (column found ignoring case) go into RW Case Notes on each record's row and never
    make a difference; without the column they are blank."""
    rows, _ = compare(tmp_path, [_POOLING.format(status='GRANTED')],
                      [rw_row('33068', 'Ordered', record_id='A') + ['Waiting on order; call Bob'],
                       rw_row('33068', '', record_id='B') + ['']], header=_RW_HEADER + [' case notes'])
    assert [(r['Record Id'], r['Differences'], r['RW Case Notes']) for r in rows] == \
        [('A', '', 'Waiting on order; call Bob'), ('B', 'Status', '')]
    [row], _ = compare(tmp_path, [_POOLING.format(status='GRANTED')], [rw_row('33068', 'Ordered')])
    assert row['RW Case Notes'] == ''


_GROUPED = ('Case 32840 (ENERPLUS): Pooling and Spacing AMENDMENT: GRANTED\n'
            '(4 WELLS)\n13-144N-95W-ALL\n16-144N-94W-ALL\n(4 WELLS)\n24-144N-95W-ALL\n19-144N-94W-ALL')  # two well groups in Slack


def test_case_is_compared_with_the_number_and_each_plain_sub_case(tmp_path):
    """A case gets one row per RW record numbered <number> or <number>-k, in number order, whatever its
    well groups; Roman-numeral sub-cases and other numbers sharing the prefix are left out."""
    rows, counts = compare(tmp_path, [_GROUPED], [
        rw_row('32840-2', 'Ordered', 'Establish Overlapping Spacing'),
        rw_row('32840-1', 'Ordered', 'Establish Overlapping Spacing'),
        rw_row('32840-3', '', 'Establish Overlapping Spacing'),
        rw_row('32840-II-1', 'Ordered', 'Establish Spacing'),
        rw_row('328401', 'Ordered', 'Establish Spacing')])
    assert [(r['Case Number'], r['RW Case Number'], r['Result']) for r in rows] == [
        ('32840', '32840-1', 'Match'), ('32840', '32840-2', 'Match'), ('32840', '32840-3', 'Differs')]
    assert counts == {'Match': 2, 'Differs': 1}


def test_roman_sub_cases_are_used_when_rw_has_no_plain_record(tmp_path):
    """A field-rules case RW only has as Roman-numeral sub-cases is compared with all of them, not Not in RW."""
    rows, _ = compare(tmp_path, ['Case 32665 (X): Redefine Field Limits: GRANTED'],
                      [rw_row('32665-V-2', 'Ordered'), rw_row('32665-V-1', '')])
    assert [(r['RW Case Number'], r['Differences']) for r in rows] == [('32665-V-1', 'Status'), ('32665-V-2', '')]


def test_every_number_of_a_case_is_looked_up(tmp_path):
    """'196-2025 / 32-2025 FED' is compared with both numbers' records; another year's 196 isn't one of them."""
    rows, _ = compare(tmp_path, ['Case 196-2025 / 32-2025 FED (CONTINENTAL): Pooling and Spacing: GRANTED'],
                      [rw_row('196-2025', 'Ordered'), rw_row('32-2025', 'Ordered'), rw_row('196-2024', 'Ordered')])
    assert [r['RW Case Number'] for r in rows] == ['196-2025', '32-2025']


def test_case_missing_from_rw(tmp_path):
    """A case RW doesn't have gets one Not in RW row with only the Slack side."""
    rows, counts = compare(tmp_path, [_POOLING.format(status='GRANTED') +
                                      '\nCase 32275 (CONTINENTAL): Pooling: GRANTED\n19-155N-97W-ALL'],
                           [rw_row('33068', 'Ordered')])
    assert [(r['Case Number'], r.get('RW Case Number', ''), r.get('Record Id', ''), r['Result']) for r in rows] == [
        ('33068', '33068', 'REC-33068', 'Match'), ('32275', '', '', 'Not in RW')]
    assert counts == {'Match': 1, 'Not in RW': 1}


def test_duplicate_rw_records_are_both_compared(tmp_path):
    """Two RW records with the same case number each get a row."""
    rows, _ = compare(tmp_path, [_POOLING.format(status='GRANTED')],
                      [rw_row('33068', 'Ordered', record_id='A'), rw_row('33068', '', record_id='B')])
    assert [(r['Record Id'], r['Result']) for r in rows] == [('A', 'Match'), ('B', 'Differs')]


def test_rw_continuation_rows_are_skipped(tmp_path):
    """A report row with a blank Record Id (an extra list value of the case above) isn't a record."""
    rows, _ = compare(tmp_path, [_POOLING.format(status='GRANTED')],
                      [rw_row('33068', 'Ordered'), ['', '', '', '', '', '9-158N-96W']])
    assert [(r['Record Id'], r['Result']) for r in rows] == [('REC-33068', 'Match')]


# ---- Upload file (Slack_Docket_Upload.csv) ----

_UPLOAD_RW_HEADER = _RW_HEADER + ['Order Status', 'Number of Wells', 'Date Order Signed',
                                  'Case Notes']  # RW report with every upload column


def urow(number, status='', order='', wells='', signed='', notes='', record_id=None):
    """One _UPLOAD_RW_HEADER report row."""
    return rw_row(number, status, record_id=record_id) + [order, wells, signed, notes]


def uploads(tmp_path, texts, rw_rows, header=_UPLOAD_RW_HEADER):
    """(upload header, upload rows as dicts, comparison rows) for messages (newest first, posted
    09/30/2026, link https://slack.example/p1) compared with an RW report."""
    path = write_csv(str(tmp_path / 'rw.csv'), [header] + rw_rows)
    rw, columns = du._read_rw(path)
    cases = du._cases_from([export_row(t, micro=f'{i:06d}') for i, t in enumerate(texts)])
    results, _ = du._compare(cases, rw, columns)
    up_header = du._upload_columns(columns)
    return up_header, [dict(zip(up_header, r['Upload'])) for r in results if r['Upload']], results


def test_upload_columns_are_named_as_in_the_rw_report():
    """The upload has Record Id, Case Status, Order Status, Number of Wells, Date Order Signed and Case
    Notes, named as the RW report names them."""
    assert du._UPLOAD_HEADER == ['Record Id', 'Case Status', 'Order Status', 'Number of Wells',
                                 'Date Order Signed', 'Case Notes']


def test_upload_row_values(tmp_path):
    """GRANTED with 3 wells against a Docketed record: Slack's values, the post date and the link + old notes."""
    header, [row], _ = uploads(tmp_path, [_POOLING.format(status='GRANTED') + '\n(3 Wells)'],
                               [urow('33068', 'Docketed', '', '2', '', 'Waiting on hearing')])
    assert header == du._UPLOAD_HEADER
    assert row == {'Record Id': 'REC-33068', 'Case Status': 'Ordered', 'Order Status': 'Granted',
                   'Number of Wells': '3', 'Date Order Signed': '09/30/2026',
                   'Case Notes': 'https://slack.example/p1 Waiting on hearing'}


def test_upload_carries_rw_values_where_slack_has_none(tmp_path):
    """With no result in Slack, Case Status, Order Status and Date Order Signed keep RW's values (a blank
    would clear Ark's); here only the well count differs and comes from Slack."""
    _, [row], _ = uploads(tmp_path, ['Case 33068 (X): Pooling:\n(3 Wells)'],
                          [urow('33068', 'Docketed', 'Pending', '2', '08/01/2026', '')])
    assert row == {'Record Id': 'REC-33068', 'Case Status': 'Docketed', 'Order Status': 'Pending',
                   'Number of Wells': '3', 'Date Order Signed': '08/01/2026', 'Case Notes': 'https://slack.example/p1'}


def test_upload_keeps_rw_well_count_when_slack_has_none(tmp_path):
    """A case with no (N Wells) line keeps RW's Number of Wells as it is."""
    _, [row], _ = uploads(tmp_path, [_POOLING.format(status='GRANTED')], [urow('33068', 'Docketed', '', '4.0')])
    assert row['Number of Wells'] == '4.0'


def test_dismissed_writes_a_blank_order_status(tmp_path):
    """DISMISSED means Case Status Dismissed and a blank Order Status, written on purpose to clear RW's."""
    _, [row], _ = uploads(tmp_path, [_POOLING.format(status='DISMISSED')], [urow('33068', 'Ordered', 'Granted')])
    assert (row['Case Status'], row['Order Status'], row['Date Order Signed']) == ('Dismissed', '', '09/30/2026')


@pytest.mark.parametrize('rw_date, written', [
    ('', '09/30/2026'),                     # blank: the post's date
    ('09/01/2026', '09/30/2026'),           # earlier: moved forward
    ('2026-09-01', '09/30/2026'),
    ('10/15/2026', '10/15/2026'),           # later than the post: never moved back
    ('10/15/2026 12:00 AM', '10/15/2026'),  # a time after the date
    ('09/30/2026', '09/30/2026'),
])
def test_date_order_signed_only_moves_forward(tmp_path, rw_date, written):
    """Date Order Signed becomes the post's date only when that is later than RW's; written MM/DD/YYYY."""
    _, [row], _ = uploads(tmp_path, [_POOLING.format(status='GRANTED')], [urow('33068', 'Docketed', '', '', rw_date)])
    assert row['Date Order Signed'] == written


def test_unreadable_rw_date_is_kept_and_noted(tmp_path):
    """An RW Date Order Signed that isn't a date is kept as it is, and the differences row says so."""
    _, [row], results = uploads(tmp_path, [_POOLING.format(status='GRANTED')],
                                [urow('33068', 'Docketed', '', '', 'pending')])
    assert row['Date Order Signed'] == 'pending'
    assert results[0]['Notes'] == 'RW Date Order Signed "pending" not read as a date: kept in the upload'


def test_link_already_in_case_notes_is_not_added_again(tmp_path):
    """Case Notes that already hold the post's link are uploaded as they are."""
    _, [row], _ = uploads(tmp_path, [_POOLING.format(status='GRANTED')],
                          [urow('33068', 'Docketed', notes='https://slack.example/p1 Waiting on hearing')])
    assert row['Case Notes'] == 'https://slack.example/p1 Waiting on hearing'


def test_only_differing_records_are_uploaded(tmp_path):
    """A matching record and a case missing from RW get no upload row."""
    _, rows, _ = uploads(tmp_path, [_POOLING.format(status='GRANTED') + '\nCase 33069 (X): Pooling: GRANTED'
                                    '\nCase 32275 (X): Pooling: GRANTED'],
                         [urow('33068', 'Ordered', 'Granted'), urow('33069', 'Docketed')])
    assert [r['Record Id'] for r in rows] == ['REC-33069']


@pytest.mark.parametrize('text, rw_rows, held', [
    (_GROUPED, [urow('32840-1', 'Docketed'), urow('32840-2', 'Docketed')], 'sub-cases'),
    (_GROUPED, [urow('32840-1', 'Docketed')], 'sub-cases'),                  # its only record is a sub-case
    (_POOLING.format(status='GRANTED'), [urow('33068', 'Docketed', record_id='A'),
                                         urow('33068', 'Docketed', record_id='B')], '2 RW records'),  # duplicates
    ('Case 196-2025 / 32-2025 FED (X): Pooling: GRANTED', [urow('196-2025', 'Docketed'), urow('32-2025', 'Docketed')],
     '2 RW records'),
])
def test_cases_with_sub_cases_are_held_back(tmp_path, text, rw_rows, held):
    """A case with sub-cases or several RW records gets no upload row; each differing record says why in the
    differences sheet, for review by hand."""
    _, rows, results = uploads(tmp_path, [text], rw_rows)
    assert rows == []
    assert results and all(r['Notes'].endswith(f'Not in Slack_Docket_Upload.csv ({held}): review by hand')
                           for r in results)


def test_upload_leaves_out_columns_the_report_lacks(tmp_path):
    """An upload column the RW report doesn't have is left out, so a blank never clears Ark's value."""
    header, [row], _ = uploads(tmp_path, [_POOLING.format(status='GRANTED') + '\n(3 Wells)'],
                               [rw_row('33068', 'Docketed') + ['2']], header=_WELLS_HEADER)
    assert header == ['Record Id', 'Case Status', 'Number of Wells']
    assert row == {'Record Id': 'REC-33068', 'Case Status': 'Ordered', 'Number of Wells': '3'}


# ---- Cases not in RW (Slack_Not_In_RW.csv) ----

def test_not_in_rw_lists_the_slack_side(tmp_path):
    """Only Not in RW cases are listed, with the Slack fields and well count (blank when none) and no
    state or sections; a case RW has is left out."""
    assert du._MISSING_HEADER == ['Case Number', 'Applicant', 'Case Type', 'Case Status', 'Order Status',
                                  'Well Count']
    rows = not_in_rw(tmp_path, [
        'NORTH DAKOTA:\n' + _POOLING.format(status='GRANTED') +
        '\nCase 32275 (CONTINENTAL): Pooling and Spacing: DENIED\n(2 Wells)\n19-155N-97W-ALL\n20-155N-97W-W2\n'
        'MONTANA:\nCase 85-2026 (KRAKEN): Spacing: DISMISSED\n1-25N-58E-ALL'],
        [rw_row('33068', 'Ordered')])
    assert rows == [
        {'Case Number': '32275', 'Applicant': 'CONTINENTAL', 'Case Type': 'Pooling and Spacing',
         'Case Status': 'Ordered', 'Order Status': 'Denied', 'Well Count': 2},
        {'Case Number': '85-2026', 'Applicant': 'KRAKEN', 'Case Type': 'Spacing',
         'Case Status': 'Dismissed', 'Order Status': '', 'Well Count': ''}]


def test_not_in_rw_one_row_per_case(tmp_path):
    """A case with several well groups is still one Not in RW row, with their wells added up."""
    assert [(r['Case Number'], r['Well Count']) for r in not_in_rw(tmp_path, [_GROUPED], [])] == [('32840', 8)]


@pytest.mark.parametrize('slack, case_status, order_status', [
    ('GRANTED', 'Ordered', 'Granted'),
    ('PARTIALLY GRANTED', 'Ordered', 'Granted'),
    ('DENIED', 'Ordered', 'Denied'),
    ('DISMISSED', 'Dismissed', ''),
    ('', '', ''),                     # no status in Slack: both left blank
])
def test_not_in_rw_status_in_rw_terms(tmp_path, slack, case_status, order_status):
    """The Slack result becomes RW's Case Status and Order Status (_EXPECTED_STATUS)."""
    [row] = not_in_rw(tmp_path, [_POOLING.format(status=slack)], [])
    assert (row['Case Status'], row['Order Status']) == (case_status, order_status)


# ---- The whole button with the comparison ----

def test_compares_messages_with_the_rw_report(slack, folders, fake_ark, msgboxes):
    """The RW report is downloaded to Slack_RW_Case_Info.csv; Slack_Docket_Differences.xlsx gets only the
    records that differ (matches and Not in RW cases left out) and Slack_Not_In_RW.csv the missing case."""
    slack.history = [msg(ts(2026, 9, 30, 15, 0, 0), _POOLING.format(status='GRANTED') + '\n(2 Wells)' +
                         '\nCase 33069 (X): Pooling: GRANTED\n(1 Well)'
                         '\nCase 32275 (CONTINENTAL): Pooling: GRANTED\n(3 Wells)\n19-155N-97W-ALL')]
    fake_ark.serve(du._REPORT_URL, [_WELLS_HEADER, rw_row('33068', '') + ['2'], rw_row('33069', 'Ordered') + ['1']],
                   filename='b1e2-uuid')
    run(folders)
    out = folders['DocketUpdates']
    assert fake_ark.downloads == [du._REPORT_URL]
    assert read_csv(out / 'Slack_RW_Case_Info.csv')[1][2] == '33068'
    assert not (out / 'Slack_Docket_Differences.csv').exists()
    rows, _, _ = read_differences(out / 'Slack_Docket_Differences.xlsx')
    assert rows[0] == du._DIFF_HEADER
    row = dict(zip(du._DIFF_HEADER, rows[1]))
    assert len(rows) == 2 and (row['Differences'], row['Case Number'], row['Slack Well Count'], row['RW Well Count']) \
        == ('Status', '33068', 2, 2)
    assert read_csv(out / 'Slack_Not_In_RW.csv') == [du._MISSING_HEADER,
                                               ['32275', 'CONTINENTAL', 'Pooling', 'Ordered', 'Granted', '3']]
    assert read_csv(out / 'Slack_Docket_Upload.csv') == [['Record Id', 'Case Status', 'Number of Wells'],
                                                   ['REC-33068', 'Ordered', '2']]
    text = _box(msgboxes)
    assert 'Saved 1 message(s)' in text and str(out / 'Parsed_Slack_Messages.csv') in text
    assert 'Compared 3 case(s) from the messages with 2 RW case(s): 1 match, 1 differ, 1 not in RW.' in text
    assert '1 RW record in the upload file; 0 differing record(s) with sub-cases left out' in text
    assert 'Nothing was uploaded' in text
    assert 'Not compared (no such column in the RW report): Order Status' in text
    assert ('Left out of the upload (no such column in the RW report): Order Status, Date Order Signed, '
            'Case Notes') in text
    for name in ('Slack_Docket_Differences.xlsx', 'Slack_Docket_Upload.csv', 'Slack_Not_In_RW.csv'):
        assert str(out / name) in text


def test_differences_workbook_highlights_the_differing_cells(tmp_path):
    """Each difference highlights its Slack and RW cells (Status: Slack Status + RW Case Status, ...);
    other cells aren't highlighted, the header is bold and frozen with a filter, and Link is clickable."""
    rows, _ = compare(tmp_path, [_POOLING.format(status='GRANTED') + '\n(2 Wells)', 'Case 33069 (X): Spacing: DENIED'],
                      [rw_row('33068', '') + ['', '3'], rw_row('33069', 'Ordered', 'Establish Spacing') + ['Granted', '']],
                      header=_RW_HEADER + ['Order Status', 'Number of Wells'])
    path = tmp_path / 'diff.xlsx'
    du._write_differences(str(path), rows)
    values, highlighted, ws = read_differences(path)
    assert [r[0] for r in values[1:]] == ['Status, Order Status, Well Count', 'Order Status']
    assert highlighted == {(1, 'Slack Status'), (1, 'RW Case Status'), (1, 'RW Order Status'),
                           (1, 'Slack Well Count'), (1, 'RW Well Count'),
                           (2, 'Slack Status'), (2, 'RW Order Status')}
    assert ws['A1'].font.b and ws.freeze_panes == 'A2' and ws.auto_filter.ref == ws.dimensions
    link = ws.cell(2, du._DIFF_HEADER.index('Link') + 1)
    assert link.hyperlink.target == 'https://slack.example/p1'


def test_no_cases_means_no_ark_download(slack, folders, fake_ark, msgboxes):
    """Messages without any case don't open Ark; the differences workbook, upload (Record Id only: no
    report columns known) and Not in RW CSV are header-only and the box says so."""
    slack.history = [msg(ts(2026, 9, 24, 13, 0, 0), 'Board Order Updates from 9/24/2026:\nNORTH DAKOTA:\nN/A')]
    run(folders)
    assert fake_ark.downloads == []
    assert read_differences(folders['DocketUpdates'] / 'Slack_Docket_Differences.xlsx')[0] == [du._DIFF_HEADER]
    assert read_csv(folders['DocketUpdates'] / 'Slack_Docket_Upload.csv') == [['Record Id']]
    assert read_csv(folders['DocketUpdates'] / 'Slack_Not_In_RW.csv') == [du._MISSING_HEADER]
    assert 'No board order cases found' in _box(msgboxes)


@pytest.mark.parametrize('locked', ['Slack_Docket_Differences.xlsx', 'Slack_Docket_Upload.csv', 'Slack_Not_In_RW.csv'])
def test_locked_comparison_file_says_close_excel(slack, folders, fake_ark, monkeypatch, msgboxes, locked):
    """If Excel has the differences workbook, the upload or the Not in RW CSV open, the box names that file."""
    slack.history = [msg(ts(2026, 9, 30, 15, 0, 0), _POOLING.format(status='GRANTED'))]
    fake_ark.serve(du._REPORT_URL, [_RW_HEADER])
    for name in ('_write_csv', '_write_differences'):
        def write(path, *args, real=getattr(du, name)):
            """The real writer, except that it fails on the locked file."""
            if path.endswith(locked):
                raise PermissionError(13, 'Permission denied', path)
            real(path, *args)
        monkeypatch.setattr(du, name, write)
    du.main()
    text = _box(msgboxes)
    assert locked in text and 'Close it in Excel' in text
