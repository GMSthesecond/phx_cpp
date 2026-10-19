"""Index button (doc_inventory.py, run through doc_inventory_roosevelt / richland / divide.py):
finds unindexed PDFs in Box.

Each county runs against a temp folder tree, the [Folders] Index folder (temp, via the folders
fixture) and a temp copy of the index workbook. Also checks the stdout tags main.cpp's progress bar
and summary depend on, and the readable stops (workbook missing or open in Excel).
"""
import ctypes
import dataclasses
import re
import runpy
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from PyPDF2 import PdfWriter

import doc_inventory
import doc_inventory_divide
import doc_inventory_richland
import doc_inventory_roosevelt
from conftest import REPO, SANDBOX
from fakes import read_csv

COUNTY_SCRIPTS = [doc_inventory_roosevelt, doc_inventory_richland, doc_inventory_divide]  # Index runs these in order


def _write_pdf(path, pages):
    """Writes a blank PDF with the given number of pages."""
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=612, height=792)
    with open(path, 'wb') as f:
        w.write(f)


@pytest.fixture(params=COUNTY_SCRIPTS, ids=lambda m: m.COUNTY.name)
def county(request, tmp_path, monkeypatch, folders):
    """One county script (.mod) pointed at a temp Box tree, the temp Index folder (.csv) and a temp
    workbook (.excel). The tree has a new 2-page PDF, an already-indexed PDF, a corrupt PDF and a non-PDF."""
    mod = request.param
    root = tmp_path / 'Box' / 'Books'
    (root / 'Book 1').mkdir(parents=True)
    (root / 'Book 2').mkdir()
    _write_pdf(root / 'Book 1' / 'new_doc.pdf', 2)
    _write_pdf(root / 'Book 1' / 'old_doc.pdf', 1)
    (root / 'Book 2' / 'broken.pdf').write_bytes(b'not really a pdf')
    (root / 'Book 2' / 'notes.txt').write_text('skip me')

    excel = tmp_path / 'Index of Scanned Docs.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = mod.COUNTY.name
    ws.append(['FileName', 'Folder', 'PageCount'])
    ws.append(['old_doc.pdf', 'Book 1', 1])
    wb.save(excel)

    monkeypatch.setattr(mod, 'COUNTY', dataclasses.replace(mod.COUNTY, roots=[str(root)]))
    monkeypatch.setattr(doc_inventory, 'workbook_path', lambda: str(excel))
    return SimpleNamespace(mod=mod, csv=folders['Index'] / f'{mod.COUNTY.name}.csv', excel=excel)


def _tags(stdout):
    """(progress tags, done counts) parsed the way main.cpp's RunOneIndexScript does:
    split on \\r or \\n, then find the tag anywhere in the line."""
    progress, done = [], []
    for line in re.split(r'[\r\n]', stdout):
        if '@PROGRESS ' in line:
            a, b = line[line.index('@PROGRESS ') + 10:].split()[:2]
            progress.append((int(a), int(b)))
        elif '@DONE ' in line:
            done.append(int(line[line.index('@DONE ') + 6:].split()[0]))
    return progress, done


def test_finds_new_pdfs_and_appends_them(county, capsys):
    """New PDFs go to the CSV and the county tab; indexed PDFs and non-PDFs are skipped."""
    county.mod.main()

    rows = read_csv(county.csv)
    assert rows[0] == ['FileName', 'Folder', 'PageCount']
    found = {r[0]: r for r in rows[1:]}
    assert set(found) == {'new_doc.pdf', 'broken.pdf'}
    assert found['new_doc.pdf'] == ['new_doc.pdf', 'Book 1', '2']
    assert found['broken.pdf'][2] in ('Unreadable', 'Encrypted/Unreadable')

    ws = load_workbook(county.excel)[county.mod.COUNTY.name]
    tab = [list(r) for r in ws.iter_rows(values_only=True)]
    assert ['new_doc.pdf', 'Book 1', 2] in tab  # page count stays numeric
    assert len(tab) == 4


def test_each_county_keeps_its_own_backup(county, capsys):
    """The workbook backup is per county, taken just before that county writes, so running all three
    leaves three backups instead of one overwritten twice."""
    county.mod.main()
    backup = county.excel.parent / f'Index of Scanned Docs.xlsx.{county.mod.COUNTY.name}.bak'
    assert backup.exists()
    with open(backup, 'rb') as f:  # openpyxl checks the extension of a path, not of a file
        tab = [list(r) for r in load_workbook(f)[county.mod.COUNTY.name].iter_rows(values_only=True)]
    assert len(tab) == 2  # the state before this county added its rows
    assert not (county.excel.parent / 'Index of Scanned Docs.xlsx.bak').exists()


def test_stdout_tags_feed_the_progress_bar(county, capsys):
    """main.cpp finds @PROGRESS (total -1 for Richland's marquee) and @DONE <new docs>."""
    county.mod.main()

    progress, done = _tags(capsys.readouterr().out)
    assert done == [2]
    assert progress, 'no @PROGRESS tag printed'
    expected_total = -1 if county.mod is doc_inventory_richland else 2
    assert all(total == expected_total for _, total in progress)


def test_second_run_adds_nothing(county, capsys):
    """Re-running right away finds no new documents (they are in the workbook now)."""
    county.mod.main()
    capsys.readouterr()

    county.mod.main()

    _, done = _tags(capsys.readouterr().out)
    assert done == [0]
    assert len(read_csv(county.csv)) == 1


def test_richland_matches_names_without_extension(county, capsys):
    """Richland's tab lists names without ".pdf", so "old_doc" there skips old_doc.pdf; the others
    need the full name."""
    wb = load_workbook(county.excel)
    wb[county.mod.COUNTY.name]['A2'] = 'OLD_DOC'
    wb.save(county.excel)

    county.mod.main()

    names = {r[0] for r in read_csv(county.csv)[1:]}
    assert ('old_doc.pdf' in names) == (county.mod is not doc_inventory_richland)


def test_csv_encoding_per_county(county, capsys):
    """Divide's CSV has a UTF-8 BOM (opens cleanly in Excel); the other two don't, as before."""
    county.mod.main()
    has_bom = county.csv.read_bytes().startswith(b'\xef\xbb\xbf')
    assert has_bom == (county.mod is doc_inventory_divide)


def test_paths_come_from_settings(folders):
    """Each county CSV goes to [Folders] Index as <County>.csv, and the workbook is [Files] IndexWorkbook."""
    for mod in COUNTY_SCRIPTS:
        assert doc_inventory.output_csv(mod.COUNTY) == str(folders['Index'] / f'{mod.COUNTY.name}.csv')
    assert doc_inventory.workbook_path() == str(SANDBOX / 'Files' / 'Index of Scanned Docs.xlsx')


def test_main_returns_the_new_document_count(county, capsys):
    """main() returns how many documents it added (Mirror's Index step adds these up)."""
    assert county.mod.main() == 2
    assert county.mod.main() == 0


def test_missing_workbook_stops_with_a_readable_message(county, monkeypatch):
    """A missing workbook stops the script with a message naming it (Mirror shows it in its report)."""
    monkeypatch.setattr(doc_inventory, 'workbook_path', lambda: str(county.excel.parent / 'gone.xlsx'))

    with pytest.raises(SystemExit) as stop:
        county.mod.main()

    assert 'Excel file not found' in str(stop.value) and 'gone.xlsx' in str(stop.value)


def _lock_like_excel(path):
    """Opens path the way Excel holds an open workbook: others may read it but not write it.
    Returns (kernel32, handle) to close."""
    kernel32 = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel32.CreateFileW.restype = ctypes.c_void_p
    kernel32.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
                                     ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
    kernel32.CloseHandle.argtypes = [ctypes.c_void_p]
    # GENERIC_READ | GENERIC_WRITE, FILE_SHARE_READ, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL
    handle = kernel32.CreateFileW(str(path), 0xC0000000, 1, None, 3, 0x80, None)
    assert handle not in (None, ctypes.c_void_p(-1).value), 'could not lock the test workbook'
    return kernel32, handle


def test_workbook_open_in_excel_stops_before_the_scan(county, capsys):
    """If Excel has the workbook open, the county stops at once with "close it" (no CSV written, no
    half-finished run), instead of failing to save after the long scan."""
    kernel32, handle = _lock_like_excel(county.excel)
    try:
        with pytest.raises(SystemExit) as stop:
            county.mod.main()
    finally:
        kernel32.CloseHandle(handle)
    assert 'Close Index of Scanned Docs.xlsx in Excel' in str(stop.value)
    assert not county.csv.exists()


def test_workbook_opened_during_the_run_says_where_the_rows_are(county, capsys, monkeypatch):
    """If saving fails because Excel opened the workbook mid-run, the stop says the rows are in the CSV."""
    def locked_save(self, path):
        raise PermissionError(13, 'Permission denied', str(path))
    import openpyxl.workbook.workbook
    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, 'save', locked_save)

    with pytest.raises(SystemExit) as stop:
        county.mod.main()

    assert 'opened in Excel during the run' in str(stop.value) and str(county.csv) in str(stop.value)
    assert len(read_csv(county.csv)) == 3  # the CSV still has the new rows


def test_button_run_shows_a_stop_in_a_box(county, msgboxes, monkeypatch):
    """Launched by the Index button, a readable stop appears as an error box with the message itself
    (main.cpp only reads the @ tags, so stderr would be invisible)."""
    monkeypatch.setattr(doc_inventory, 'workbook_path', lambda: str(county.excel.parent / 'gone.xlsx'))

    with pytest.raises(RuntimeError):
        doc_inventory.run_button(county.mod.COUNTY)

    [(title, text)] = msgboxes
    assert title == f'Index ({county.mod.COUNTY.name})'
    assert 'Excel file not found' in text and 'gone.xlsx' in text


@pytest.mark.parametrize('mod', COUNTY_SCRIPTS, ids=lambda m: m.COUNTY.name)
def test_wrapper_scripts_run_their_county_as_a_button(mod, monkeypatch):
    """Each county script, run as main.cpp runs it, calls doc_inventory.run_button with its own county."""
    seen = []
    monkeypatch.setattr(doc_inventory, 'run_button', lambda c: seen.append(c.name))
    runpy.run_path(str(REPO / f'{mod.__name__}.py'), run_name='__main__')
    assert seen == [mod.COUNTY.name]
