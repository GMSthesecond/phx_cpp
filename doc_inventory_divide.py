import os
import csv
import shutil
import sys
import time
from typing import Optional, List, Set, Iterable, Tuple
import unicodedata

ROOTS = [
    r"C:\Cloud\Box\Land\Title Folder\Document Images\North Dakota\Divide County Images\Deeds",
    r"C:\Cloud\Box\Land\Title Folder\Document Images\North Dakota\Divide County Images\Misc",
    r"C:\Cloud\Box\Land\Title Folder\Document Images\North Dakota\Divide County Images\Modern",
    r"C:\Cloud\Box\Land\Title Folder\Document Images\North Dakota\Divide County Images\Mort",
    r"C:\Cloud\Box\Land\Title Folder\Document Images\North Dakota\Divide County Images\Plat",
    r"C:\Cloud\Box\Land\Title Folder\Document Images\North Dakota\Divide County Images\Probates and Judgements",
    r"C:\Cloud\Box\Land\Title Folder\Document Images\North Dakota\Divide County Images\Unnamed and Speciality Books",

]

def clean_name(name: str) -> str:
    s = unicodedata.normalize("NFC", name)
    s = s.replace("\u200b", "").replace("\ufeff", "")
    return s

OUTPUT_CSV = r"C:\Users\Ethan Mesecher\Desktop\Steph Suko\Divide.csv"

# Excel (names to skip)
EXCEL_PATH = r"C:\Cloud\Box\Ethan Mesecher\Index of Scanned Docs.xlsx"
EXCEL_SHEET = "Divide"
EXCEL_COLUMN_INDEX = 1  # Column A = 1 (1-based)

# What to put in the Folder column:
#   "parent" -> the immediate parent directory of the file
#   "root"   -> the configured root label (e.g., "A Books")
FOLDER_NAME_MODE = "parent"  # or "root"

# Only attempt page count for PDFs
PDF_ONLY = True
# ==============================================================================

# -------- Long-path helper ----------------------------------------------------
def win_long(path: str) -> str:
    if os.name == "nt":
        path = os.path.abspath(path)
        if not path.startswith("\\\\?\\"):
            if path.startswith("\\\\"):  # UNC path
                return "\\\\?\\UNC\\" + path[2:]
            else:
                return "\\\\?\\" + path
    return path

def win_strip_long(path: str) -> str:
    """Remove the long-path prefix for prettier display / CSV output."""
    if path.startswith("\\\\?\\UNC\\"):
        return "\\" + path[7:]  # -> \\server\share\...
    elif path.startswith("\\\\?\\"):
        return path[4:]
    return path

# -------- Load already-indexed filenames from Excel ---------------------------
def load_indexed_names(
    excel_path: str,
    sheet: str = "Richland",
    col_index: int = 1,  # 1-based (A=1)
    skip_header: bool = False
) -> Set[str]:
    """
    Reads the specified column (1-based) from the given sheet and returns a set
    of normalized names (lower-cased, trimmed). Assumes entries already include
    the final form you want to match (e.g., 'filename.pdf').
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("This script requires 'openpyxl'. Install with: pip install openpyxl")
        sys.exit(1)

    if not os.path.exists(excel_path):
        print(f"Excel file not found: {excel_path}")
        sys.exit(1)

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        print(f"Sheet '{sheet}' not found in workbook. Found: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[sheet]
    names: Set[str] = set()

    row_iter = ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=col_index,
        max_col=col_index,
        values_only=True
    )

    first = True
    for (val,) in row_iter:
        if val is None:
            continue
        if skip_header and first:
            first = False
            continue
        s = str(val).strip()
        if not s:
            continue
        names.add(s.lower())

    wb.close()
    return names

# -------- Fast file counting (with filter) ------------------------------------
def iter_files(root: str) -> Iterable[Tuple[str, str]]:
    """
    Yield (dirpath, filename) for files under 'root', handling permission errors.
    """
    for dirpath, dirnames, filenames in os.walk(root, onerror=lambda e: None):
        for fname in filenames:
            yield dirpath, fname

def count_files_to_process(roots: List[str], indexed_names: Set[str]) -> int:
    """
    Counts files across all roots that are NOT in the indexed_names set.
    If PDF_ONLY is True, only counts *.pdf files.
    """
    total = 0
    stack: List[str] = list(roots)
    while stack:
        current = stack.pop()
        try:
            with os.scandir(current) as it:
                for entry in it:
                    try:
                        if entry.is_dir(follow_symlinks=False):
                            stack.append(entry.path)
                        elif entry.is_file(follow_symlinks=False):
                            if PDF_ONLY and not entry.name.lower().endswith(".pdf"):
                                continue
                            fname_lower = entry.name.lower()
                            if fname_lower not in indexed_names:
                                total += 1
                    except (PermissionError, OSError):
                        continue
        except (PermissionError, OSError):
            continue
    return total

# -------- Lightweight progress printer ----------------------------------------
class Progress:
    def __init__(self, total: int, min_interval: float = 0.5):
        self.total = total
        self.done = 0
        self.last_t = 0.0
        self.min_interval = min_interval

    def tick(self, n: int = 1):
        self.done += n
        now = time.time()
        if now - self.last_t >= self.min_interval:
            self.last_t = now
            pct = (self.done / self.total * 100) if self.total > 0 else 0.0
            print(f"\rProcessing: {self.done:,}/{self.total:,} ({pct:5.1f}%)", end="", flush=True)
            print(f"@PROGRESS {self.done} {self.total}", flush=True)

    def finish(self):
        if self.total > 0:
            print(f"\rProcessing: {self.total:,}/{self.total:,} (100.0%)")
        else:
            print("\rProcessing: complete")

# -------- PDF page counter (PyPDF2) -------------------------------------------
def get_pdf_page_count(pdf_path: str) -> Optional[str]:
    try:
        from PyPDF2 import PdfReader
        with open(pdf_path, "rb") as f:
            reader = PdfReader(f)
            # Try to handle encrypted-but-readable PDFs
            try:
                _ = len(reader.pages)
            except Exception:
                # Attempt empty-password decrypt for some PDFs
                try:
                    reader.decrypt("")
                except Exception:
                    return "Encrypted/Unreadable"
            try:
                return str(len(reader.pages))
            except Exception:
                return "Encrypted/Unreadable"
    except FileNotFoundError:
        return None  # file disappeared mid-scan; leave blank
    except PermissionError:
        return "No Permission"
    except Exception:
        return "Unreadable"

# -------- Helpers for folder labeling -----------------------------------------
def folder_label_for_path(full_path: str, root_label: str) -> str:
    """
    Returns the value to write in the Folder column based on FOLDER_NAME_MODE.
    - 'parent': the immediate parent directory name (e.g., 'Book 12' or 'A Books')
    - 'root'  : the human-friendly root label passed in (e.g., 'A Books')
    """
    if FOLDER_NAME_MODE == "root":
        return root_label
    # default: parent folder
    parent = os.path.basename(os.path.dirname(full_path.rstrip("\\/")))
    return parent or root_label

# -------- Append newly-found rows to the bottom of the Excel tab --------------
def append_rows_to_excel(excel_path: str, sheet: str, rows: List[List[str]]) -> None:
    """
    Appends the same rows written to the CSV to the bottom of the matching sheet,
    so the workbook stays in sync with what was just exported. A .bak copy of the
    workbook is made first so a failed/interrupted write can be recovered from.
    """
    if not rows:
        print("No new rows to add to the Excel index.")
        return

    from openpyxl import load_workbook

    backup_path = excel_path + ".bak"
    shutil.copy2(excel_path, backup_path)

    wb = load_workbook(excel_path)
    if sheet not in wb.sheetnames:
        print(f"Sheet '{sheet}' not found in workbook; skipping Excel update.")
        wb.close()
        return

    ws = wb[sheet]
    for fname, folder_val, page_count in rows:
        # Keep PageCount numeric (matches existing rows) unless it's a non-numeric
        # status like "Encrypted/Unreadable"
        page_count_val = int(page_count) if str(page_count).isdigit() else page_count
        ws.append([fname, folder_val, page_count_val])
    wb.save(excel_path)
    wb.close()
    print(f"Added {len(rows):,} row(s) to the '{sheet}' tab in {excel_path}")

# -------- Main loop with Excel-driven filtering --------------------------------
def main():
    # Prepare roots with long-path handling and a simple label
    prepared_roots: List[Tuple[str, str]] = []
    for r in ROOTS:
        long_r = win_long(r)
        label = os.path.basename(r.rstrip("\\/")) or r  # e.g., 'A Books'
        prepared_roots.append((long_r, label))

    # Load the index of filenames we should SKIP (exact name matching, case-insensitive)
    indexed_names = load_indexed_names(
        EXCEL_PATH,
        sheet=EXCEL_SHEET,
        col_index=EXCEL_COLUMN_INDEX,
    )

    # Count how many files we will actually process (for progress)
    total_files = count_files_to_process([lr for (lr, _) in prepared_roots], indexed_names)
    prog = Progress(total_files, min_interval=0.5)

    processed = 0
    skipped = 0
    new_rows: List[List[str]] = []

    # Write only files not already listed in Excel
    with open(OUTPUT_CSV, mode="w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(
            f,
            dialect="excel",  # Excel-friendly defaults
            quoting=csv.QUOTE_MINIMAL,  # only quote fields that need it
            doublequote=True,  # escape internal " as ""
            escapechar="\\",  # extra safety if needed
            lineterminator="\n"  # consistent newlines
        )
        # NOTE: Folder column added between FileName and PageCount
        writer.writerow(["FileName", "Folder", "PageCount"])

        for long_root, root_label in prepared_roots:
            for dirpath, fname in iter_files(long_root):
                fname_lower = fname.lower()

                # Skip non-PDF if PDF_ONLY is True
                if PDF_ONLY and not fname_lower.endswith(".pdf"):
                    continue

                # Skip if already indexed (exact match on filename only)
                if fname_lower in indexed_names:
                    skipped += 1
                    continue

                full_path = os.path.join(dirpath, fname)

                # Compute folder label
                folder_val = folder_label_for_path(full_path, root_label)

                page_count = ""
                try:
                    # Count pages only for PDFs (or for all if PDF_ONLY is False but still PDF)
                    if fname_lower.endswith(".pdf"):
                        page_count = get_pdf_page_count(full_path) or ""
                    # Write CSV row
                    row = [fname, folder_val, page_count]
                    writer.writerow(row)
                    new_rows.append(row)
                    processed += 1
                except Exception:
                    # Skip any single-file failure and continue
                    pass
                finally:
                    # Progress tick only for files we actually processed/attempted
                    prog.tick(1)

    prog.finish()
    print(f"Processed: {processed:,}  |  Skipped (already indexed): {skipped:,}")
    print(f"@DONE {processed}", flush=True)
    print(f"Report written to: {OUTPUT_CSV}")

    append_rows_to_excel(EXCEL_PATH, EXCEL_SHEET, new_rows)

if __name__ == "__main__":
    main()
