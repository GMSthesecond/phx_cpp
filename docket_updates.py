"""Docket Updates, launched by the "Docket Updates" button.

Reads the ND and MT sheets of the Hearing Dockets workbook in Box ([Files] HearingDockets, default
_DEFAULT_WORKBOOK), downloads the RW cases report from Ark (RW_Case_Info.csv) and compares every case
that matters (an Operator with "Phoenix" in it, or already in RW) with its RW record(s): Case Status,
Order Status, Case Type, Sections, Hearing Date, Order Number and Date Order Signed.
The UT, WY and Sheet2 sheets are not read yet (see _SHEETS and docs/KNOWN_ISSUES.md).

Writes to [Folders] DocketUpdates (default: Desktop\\Docket Updates), replacing the last run:
- Docket_Differences.csv: one row per workbook case / RW record pairing (Match / Differs), Phoenix
  cases missing from RW (Not in RW) and RW cases missing from the workbook (Not in workbook).
- Docket_Upload.csv: Record Id plus every compared field for each RW record that differs (dates as
  MM/DD/YYYY), ready for the Ark Data Loader as a Cases Update. Not uploaded automatically: load it by hand once checked.
- Not_In_RW.csv: the workbook side of every Not in RW row, in RW's column names (for an Insert).
Shows a message box with the counts when done. Read-only: nothing is changed in Box or Ark.
"""
import csv
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime

import openpyxl

import ark_common

_DEFAULT_DIR = r'C:\Users\Ethan Mesecher\Desktop\Docket Updates'  # output folder used when [Folders] DocketUpdates is unset
_DEFAULT_WORKBOOK = (r'C:\Cloud\Box\Land\01. Permit & Hearing Tracking\02. State Permits & Hearings'
                     r'\ND,MT,UT,WY - Hearing Dockets.xlsx')  # workbook used when [Files] HearingDockets is unset
# Sheets compared. UT (mostly no case numbers, lot descriptions), WY (no Case Type / Order Status) and
# Sheet2 (state website links) are left out for now; they need their own rules before they're added.
_SHEETS = ('ND', 'MT')
_WB_COLUMNS = {  # field -> workbook header (ND and MT use the same headers)
    'number': 'Case No.', 'type': 'Case Type', 'operator': 'Operator', 'sections': 'Proposed Spacing Unit',
    'hearing': 'Hearing Date', 'state': 'State', 'county': 'County', 'status': 'Order Status',
    'order_number': 'Order Number', 'signed': 'Date Order Signed'}
_OPTIONAL_WB = {'state', 'county'}  # workbook columns that may be missing without stopping the run
_PHOENIX = 'phoenix'  # an Operator containing this (any case) puts a case in scope even if RW lacks it

_REPORT_URL = 'https://ark.phoenixenergy.com/report?recordId=d240bce207ca4cac3da5f8dc'  # Ark report of RW cases
_REPORT_FILE = 'RW_Case_Info.csv'  # saved name of the RW cases report (overwritten each run)
_FIELDS = ('Case Status', 'Order Status', 'Case Type', 'Sections', 'Hearing Date', 'Order Number',
           'Date Order Signed')  # compared fields, named as RW report columns and Data Loader fields
_DATE_FIELDS = {'Hearing Date', 'Date Order Signed'}  # compared as dates, uploaded as MM/DD/YYYY
_REQUIRED_RW = ('Record Id', 'Case Number')  # report columns the comparison can't run without
_UTAH_NUMBER_RE = re.compile(r'^\d{4}-\d{3}')  # Utah case numbers (2026-027-15): skipped while UT isn't read

_DIFF_FILE = 'Docket_Differences.csv'  # comparison file name (overwritten each run)
_DIFF_HEADER = (['Result', 'Differences', 'Case Number', 'RW Case Number', 'Record Id', 'Sheet', 'Rows',
                 'State', 'County', 'Group', 'Workbook Operator', 'RW Applicant']
                + [f'{side} {f}' for f in _FIELDS for side in ('Workbook', 'RW')]
                + ['Missing from RW', 'Extra in RW', 'Notes'])  # comparison columns
_MATCH, _DIFFERS, _NOT_IN_RW, _NOT_IN_WORKBOOK = 'Match', 'Differs', 'Not in RW', 'Not in workbook'  # Result values
_UPLOAD_FILE = 'Docket_Upload.csv'  # Cases Update file for the Data Loader (overwritten each run)
_MISSING_FILE = 'Not_In_RW.csv'  # workbook side of every Not in RW row (overwritten each run)
_MISSING_HEADER = ['Case Number', 'State', 'Applicant'] + list(_FIELDS)  # Not_In_RW.csv columns (RW names)

_STATUS_RE = re.compile(r'PARTIALLY GRANTED|GRANTED|DISMISSED|DIMISSED|DENIED|WITHDRAWN')  # results found in Order Status (DIMISSED is a typo in the sheet)
_EXPECTED_STATUS = {  # workbook result -> (RW Case Status, RW Order Status); a blank Order Status isn't compared
    'GRANTED': ('Ordered', 'Granted'),
    'PARTIALLY GRANTED': ('Ordered', 'Granted'),
    'DENIED': ('Ordered', 'Denied'),
    'DISMISSED': ('Dismissed', ''),
    'DIMISSED': ('Dismissed', ''),
    'WITHDRAWN': ('Dismissed', ''),
}
_DOCKETED = ('Docketed', '')  # expected for a case with no result and a hearing today or later
_FILL_DOWN_CASES = 50  # an Order Number on more cases than this is a fill-down slip (35278 is on ~2,600 ND rows), so ignored
_PREFIX_RE = re.compile(r'\(\s*continued?\s*\)', re.I)  # "(Continued)" in front of a re-listed case's type
_TYPE_MAP = {  # workbook Case Type (uppercase, "(Continued)" removed) -> RW Case Type; others aren't compared
    'POOLING': 'Pooling',
    'ESTABLISH SPACING': 'Establish Spacing',
    'ESTABLISH OVERLAPPING SPACING': 'Establish Overlapping Spacing',
    'ESTABLISH TEMPORARY SPACING': 'Establish Temporary Spacing',
    'TEMPORARY SPACING': 'Establish Temporary Spacing',
    'ESTABLISH OVERLAPPING TEMPORARY SPACING': 'Establish Overlapping Temporary Spacing',
    'ESTABLISH PERMANENT SPACING': 'Establish Permanent Spacing',
    'PERMANENT SPACING': 'Establish Permanent Spacing',
    'INCREASED DENSITY': 'Increased Density',
    'INCREASE DENSITY': 'Increased Density',
    'ADDITIONAL WELLS': 'Increased Density',
    'AUTHORIZE WELLS': 'Increased Density',
    'AUTHORIZE WELL': 'Increased Density',
    'PERMIT TO DRILL': 'Permit to Drill',
    'AUTHORIZE DRILLING': 'Permit to Drill',
    'COMMINGLING': 'Commingling',
    'COMMINGLE': 'Commingling',
    'SWD': 'SWD Well Permit',
    'SWD WELL PERMIT': 'SWD Well Permit',
    'SWD WELL CONVERSION': 'SWD Well Conversion',
    'CONVERSION TO SALTWATER DISPOSAL': 'SWD Well Conversion',
    'SHOW CAUSE': 'Show Cause',
    'REVOKE PERMIT': 'Show Cause',
    'REVOKE PERMITS': 'Show Cause',
    'REVOKE PEMIT': 'Show Cause',
    'REVOKING PERMITS': 'Show Cause',
    'EXTEND DRILLING DEADLINE': 'Extend Drilling Deadline',
    'PERMIT PROTEST': 'Permit Protest',
    'PERMIT/PROTEST': 'Permit Protest',
    'RECOMPLETION CERTIFICATION': 'Recompletion Certification',
}
_SPACING_TYPES = ('Establish Spacing', 'Establish Overlapping Spacing', 'Establish Temporary Spacing',
                  'Establish Overlapping Temporary Spacing', 'Establish Permanent Spacing')  # RW spacing types
_TYPE_FAMILIES = {  # workbook type -> RW types that also count as a match (the sheet writes the generic
    'Establish Spacing': _SPACING_TYPES,  # ESTABLISH SPACING for cases RW records as overlapping, permanent, ...)
}
_WELL_TYPES = {'SWD Well Permit', 'SWD Well Conversion', 'Commingling', 'Show Cause', 'Permit Protest',
               'Recompletion Certification'}  # the sheet gives the well's quarter-quarter, RW the section: compared by whole section
_WB_SECTION_RE = re.compile(r'^(\d{1,2})-(\d{1,3})([NS])-(\d{1,3})([EW])(?:-(.*))?$', re.I)  # workbook line: S-TN-RW-PART (no spaces)
_RW_SECTION_RE = re.compile(r'^(?:([NSEW0-9/]+)-)?(\d{1,2})-(\d{1,3})([NS])-(\d{1,3})([EW])$')  # RW section: optional part prefix (NW-) then S-TN-RW
_PART_RE = re.compile(r'[NS][EW]|[NSEW]2')  # one aliquot step: a quarter (NE) or a half (W2)
_QUARTERS = (('NW', 0, 0), ('NE', 0, 2), ('SW', 2, 0), ('SE', 2, 2))  # quarter name and its top-left cell in the 4x4 grid


class InputError(Exception):
    """The workbook or the RW report can't be used (missing, locked, or a column is gone); the text
    is shown to the user as is."""


def _notify(message, flags=0x40):
    """Show a message box titled "Docket Updates" (info icon unless flags say otherwise)."""
    ark_common.notify('Docket Updates', message, flags)


def _write_csv(path, header, rows):
    """Write the header and rows as UTF-8 with a BOM, so Excel shows em dashes and accents."""
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)


# ---- Cell values ----

def _text(value):
    """A cell as trimmed text: None and the sheet's non-breaking-space placeholders become ''."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ' '.join(str(value).replace('\xa0', ' ').split())


def _date(value):
    """A date from a workbook cell or report text, or None if blank or unreadable."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value).split(' ')[0]
    try:
        return ark_common.parse_date(text) if text else None
    except ValueError:
        return None


def _fmt_date(d):
    """MM/DD/YYYY, the way the other uploads write dates; '' for None."""
    return d.strftime('%m/%d/%Y') if d else ''


def _order_number(value):
    """Order Number as text ('35278', '238-2026'); '' for blanks and the odd date typed into the column."""
    return '' if isinstance(value, (datetime, date)) else _text(value)


# ---- Sections: each section is a 4x4 grid of quarter-quarter cells (row 0 north, column 0 west) ----

def _part_cells(part):
    """Grid cells (row, col) an aliquot covers: ALL / AL / blank, W2, NE, S2SE, E2W2, ...
    Steps apply right to left (S2SE is the south half of the SE quarter). None if unreadable (lots)."""
    part = part.upper().replace('/', '').replace(' ', '')
    if part in ('', 'ALL', 'AL'):
        return {(r, c) for r in range(4) for c in range(4)}
    steps = _PART_RE.findall(part)
    if ''.join(steps) != part:
        return None
    r0, r1, c0, c1 = 0, 4, 0, 4
    for step in reversed(steps):
        for ch in step:
            if ch in 'NS' and r1 - r0 > 1:
                mid = (r0 + r1) // 2
                r0, r1 = (r0, mid) if ch == 'N' else (mid, r1)
            elif ch in 'EW' and c1 - c0 > 1:
                mid = (c0 + c1) // 2
                c0, c1 = (c0, mid) if ch == 'W' else (mid, c1)
    return {(r, c) for r in range(r0, r1) for c in range(c0, c1)}


def _section_key(sec, twp, tdir, rng, rdir):
    """Section name the way RW writes it, e.g. 11-156N-91W."""
    return f'{int(sec)}-{int(twp)}{tdir.upper()}-{int(rng)}{rdir.upper()}'


def _workbook_cells(text):
    """(cells, unread lines) of a Proposed Spacing Unit cell: one section per line, such as
    24-160N-90W-ALL, 6-146N-99W (no part means all) or 4-23N-59E-W2E2, E2W2. Cells are
    (section, row, col). Lines that can't be read (lots, a field name) are returned as written."""
    cells, unread = set(), []
    for raw in str(text or '').replace('\xa0', ' ').split('\n'):
        line = ''.join(raw.split())
        if not line:
            continue
        m = _WB_SECTION_RE.match(line)
        parts = [_part_cells(p) for p in (m[6] or '').split(',')] if m else [None]
        if any(p is None for p in parts):
            unread.append(raw.strip())
            continue
        key = _section_key(*m.groups()[:5])
        for got in parts:
            cells |= {(key, r, c) for r, c in got}
    return cells, unread


def _rw_cells(sections):
    """Cells of an RW Sections value ("NW-11-156N-91W, 12-156N-91W"). An unreadable entry becomes
    its own cell (text, None, None), so it always shows as a difference."""
    cells = set()
    for item in (sections or '').split(','):
        item = ''.join(item.split()).upper()
        if not item:
            continue
        m = _RW_SECTION_RE.match(item)
        got = _part_cells(m[1] or 'ALL') if m else None
        if got is None:
            cells.add((item, None, None))
            continue
        key = _section_key(*m.groups()[1:])
        cells |= {(key, r, c) for r, c in got}
    return cells


def _section_order(key):
    """Sort key for section names: township, range, then section; unreadable text last."""
    nums = [int(n) for n in re.findall(r'\d+', key)]
    return (0, nums[1], nums[2], nums[0]) if len(nums) == 3 else (1, 0, 0, 0)


def _describe(cells):
    """RW-style text for cells: 11-156N-91W for a whole section, NW-11-156N-91W for a whole quarter,
    NENW-11-156N-91W for a quarter-quarter, sorted by township, range, section."""
    by_key = defaultdict(set)
    for key, r, c in cells:
        by_key[key].add((r, c))
    out = []
    for key in sorted(by_key, key=lambda k: (_section_order(k), k)):
        got = by_key[key]
        if (None, None) in got or len(got) == 16:
            out.append(key)
            continue
        for name, r0, c0 in _QUARTERS:
            quarter = {(r0 + dr, c0 + dc) for dr in (0, 1) for dc in (0, 1)}
            if quarter <= got:
                out.append(f'{name}-{key}')
            else:
                out += [f'{"NS"[r - r0]}{"WE"[c - c0]}{name}-{key}' for r, c in sorted(quarter & got)]
    return ', '.join(out)


def _whole_sections(cells):
    """Every cell of each section cells touches (NWNE-11-156N-91W becomes all of 11-156N-91W)."""
    return {(key, r, c) for key in {k for k, _, _ in cells} for r in range(4) for c in range(4)}


# ---- Reading the workbook ----

def _status(text):
    """Result in an Order Status cell (GRANTED, PARTIALLY GRANTED, DENIED, DISMISSED, DIMISSED,
    WITHDRAWN) or '' (blank, CONTINUED, X, P&A, ...)."""
    m = _STATUS_RE.search(text.upper())
    return m[0] if m else ''


def _rw_type(text):
    """RW Case Type for a workbook Case Type ("(Continued) ESTABLISH SPACING" -> Establish Spacing),
    or None when _TYPE_MAP has no match (Define Field Limits, Amendment-..., ...)."""
    return _TYPE_MAP.get(' '.join(_PREFIX_RE.sub('', text).split()).upper())


def _read_sheet(ws):
    """Rows of one sheet as dicts (the _WB_COLUMNS fields plus sheet and row, the Excel row number).
    Rows without a case number are skipped: they can't be matched to RW. Raises InputError when a
    needed header is missing."""
    rows = ws.iter_rows(values_only=True)
    header = [_text(h) for h in next(rows, ())]
    cols = {}
    for field, name in _WB_COLUMNS.items():
        if name in header:
            cols[field] = header.index(name)
        elif field not in _OPTIONAL_WB:
            raise InputError(f'The {ws.title} sheet of the Hearing Dockets workbook has no "{name}" column.')
    out = []
    for n, values in enumerate(rows, 2):
        def get(field):
            """This row's cell for a _WB_COLUMNS field (None if the sheet lacks the column)."""
            return values[cols[field]] if field in cols and cols[field] < len(values) else None
        number = _text(get('number'))
        if not number:
            continue
        status_text = _text(get('status'))
        out.append({'sheet': ws.title, 'row': n, 'number': number, 'type': _text(get('type')),
                    'operator': _text(get('operator')), 'sections': get('sections'),
                    'hearing': _date(get('hearing')), 'state': _text(get('state')) or ws.title,
                    'county': _text(get('county')), 'status_text': status_text, 'status': _status(status_text),
                    'order_number': _order_number(get('order_number')), 'signed': _date(get('signed'))})
    return out


def _read_workbook(path):
    """Every case-numbered row of the _SHEETS sheets. Raises InputError when the workbook is
    missing, locked or lacks a sheet / column."""
    if not os.path.isfile(path):
        raise InputError(f'Hearing Dockets workbook not found:\n{path}\n\n'
                         'Set its path in Settings (Hearing Dockets workbook).')
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise InputError(f'Could not open the Hearing Dockets workbook:\n{path}\n\n'
                         'If Box is still syncing it, wait a moment and click Docket Updates again.')
    try:
        rows = []
        for name in _SHEETS:
            if name not in wb.sheetnames:
                raise InputError(f'The Hearing Dockets workbook has no {name} sheet.')
            rows += _read_sheet(wb[name])
        return rows
    finally:
        wb.close()


def _group(rows):
    """One group (sub-case) of a case: its sheet row number(s), sections as written (text), cells,
    unreadable lines and types (the RW types of its rows' Case Types; several when one case lists
    two reliefs on the same land)."""
    cells, unread = set(), []
    for r in rows:
        got, bad = _workbook_cells(r['sections'])
        cells |= got
        unread += [b for b in bad if b not in unread]
    texts = [' '.join(str(r['sections'] or '').split()) for r in rows if _text(r['sections'])]
    return {'rows': [r['row'] for r in rows], 'text': ', '.join(dict.fromkeys(texts)),
            'cells': cells, 'unread': unread,
            'types': [t for t in dict.fromkeys(_rw_type(r['type']) for r in rows) if t]}


def _groups(rows):
    """Groups (sub-cases) of one hearing's rows, in sheet order. Rows with the same land are one group:
    the sheet lists a case once per relief (ND 32346: Authorize Drilling and AUTHORIZE WELLS on the
    same sections). Rows with no sections at all stay apart."""
    by_land = {}
    for r in rows:
        cells, unread = _workbook_cells(r['sections'])
        key = (frozenset(cells), tuple(unread)) if cells or unread else ('row', r['row'])
        by_land.setdefault(key, []).append(r)
    return [_group(same) for same in by_land.values()]


def _fill_down_numbers(rows):
    """Order Numbers shared by more than _FILL_DOWN_CASES different cases: almost certainly filled
    down by mistake, so _cases ignores them (and the Date Order Signed beside them)."""
    cases = defaultdict(set)
    for r in rows:
        if r['order_number']:
            cases[r['order_number']].add((r['sheet'], r['number']))
    return {n for n, c in cases.items() if len(c) > _FILL_DOWN_CASES}


def _cases(rows, today):
    """Workbook cases, one per sheet + case number, in sheet order. A case is re-listed for every
    hearing it is continued to, so which row supplies each field is decided here:
    - groups (sub-cases): the rows of the latest hearing date that lists any sections (all rows of
      the latest hearing if none do), in sheet order; rows on one date with different land are
      sub-cases, rows with the same land one group (_groups).
    - hearing: the latest Hearing Date. Case Type: from the latest row that has one.
    - status, Order Number, Date Order Signed: from the latest row with a recognized result (first in
      sheet order on that date), even when a later hearing re-lists the case (a note says so).
      An Order Number in _fill_down_numbers is ignored together with its Date Order Signed.
    - no recognized result anywhere: Docketed if the latest hearing is today or later; otherwise no
      status is compared (a note gives the hearing and any status text such as CONTINUED).
    Applicant (operator) is from the latest row, shown but never compared."""
    suspect = _fill_down_numbers(rows)
    by_case = {}
    for r in rows:
        by_case.setdefault((r['sheet'], r['number']), []).append(r)
    cases = []
    for (sheet, number), listed in by_case.items():
        order = sorted(listed, key=lambda r: r['hearing'] or date.min, reverse=True)  # stable: sheet order within a date
        latest = order[0]
        notes = []
        with_sections = [r for r in order if _text(r['sections'])]
        group_date = (with_sections or order)[0]['hearing']
        group_rows = sorted((r for r in listed if r['hearing'] == group_date), key=lambda r: r['row'])
        groups = _groups(group_rows)
        for g in groups:
            if g['unread']:
                notes.append('Unread sections: ' + '; '.join(g['unread']))
            if len(g['rows']) > 1:
                notes.append(f'Rows {", ".join(map(str, g["rows"]))} list the same land: one sub-case'
                             + (f' ({", ".join(g["types"])})' if len(g['types']) > 1 else ''))
        type_text = next((r['type'] for r in order if r['type']), '')
        status_row = next((r for r in order if r['status']), None)
        expected, order_number, signed = None, '', None
        if status_row:
            expected = _EXPECTED_STATUS[status_row['status']]
            order_number, signed = status_row['order_number'], status_row['signed']
            if status_row['hearing'] != latest['hearing']:
                notes.append(f'{status_row["status_text"]} is from the {_fmt_date(status_row["hearing"])} hearing; '
                             f'listed again for {_fmt_date(latest["hearing"])}')
            if order_number in suspect:
                notes.append(f'Order Number {order_number} ignored (on more than {_FILL_DOWN_CASES} cases)')
                order_number, signed = '', None
        elif latest['hearing'] and latest['hearing'] >= today:
            expected = _DOCKETED
        else:
            odd = next((r['status_text'] for r in order if r['status_text']), '')
            notes.append(f'No order status after the {_fmt_date(latest["hearing"]) or "undated"} hearing'
                         + (f' (Order Status: {odd})' if odd else ''))
        cases.append({'sheet': sheet, 'number': number, 'state': latest['state'], 'county': latest['county'],
                      'operator': latest['operator'], 'type_text': type_text, 'type': _rw_type(type_text),
                      'groups': groups, 'hearing': latest['hearing'], 'expected': expected,
                      'status_text': status_row['status_text'] if status_row else '',
                      'order_number': order_number, 'signed': signed, 'notes': notes,
                      'phoenix': any(_PHOENIX in r['operator'].lower() for r in listed)})
    return cases


# ---- Comparing with RW ----

def _read_rw(path):
    """RW cases from the Ark report as dicts of its columns plus cells (from Sections), and the
    _FIELDS columns the report has. A row with a blank Record Id continues the case above (Ark
    puts extra list values there): its Sections are added to that case. Raises InputError when
    Record Id or Case Number is missing."""
    cases = []
    with open(path, newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        columns = [c.strip() for c in reader.fieldnames or []]
        missing = [c for c in _REQUIRED_RW if c not in columns]
        if missing:
            raise InputError(f'The RW cases report has no {" or ".join(missing)} column.\n\n'
                             f'Columns found: {", ".join(columns) or "none"}')
        for row in reader:
            row = {k.strip(): (v or '').strip() for k, v in row.items() if isinstance(k, str)}
            if row.get('Record Id'):
                cases.append(row)
            elif cases and row.get('Sections'):
                cases[-1]['Sections'] = ', '.join(filter(None, [cases[-1].get('Sections'), row['Sections']]))
    for case in cases:
        case['cells'] = _rw_cells(case.get('Sections'))
    return cases, [f for f in _FIELDS if f in columns]


def _candidates(case, index):
    """RW records for a workbook case: its number itself and its sub-cases (<number>-...).
    RW has no state column, so ND and MT share the index (their number formats don't overlap)."""
    n = case['number']
    return [r for r in index.get(n.split('-')[0], [])
            if r.get('Case Number') == n or r.get('Case Number', '').startswith(n + '-')]


def _pair(case, candidates):
    """Pair each group with RW records: first every unused record with the same land, then by
    position (group k with sub-case <number>-k, a lone group with <number> itself), then a single
    group still unpaired takes every record left. Returns (pairs, unpaired groups, unpaired records)."""
    groups = case['groups']
    pairs, used, open_groups = [], [], []
    for g in groups:
        same = [r for r in candidates if r not in used and g['cells'] and r['cells'] == g['cells']]
        pairs += [(g, r) for r in same]
        used += same
        if not same:
            open_groups.append(g)
    unpaired = []
    for g in open_groups:
        k = groups.index(g) + 1
        names = {f'{case["number"]}-{k}'} if len(groups) > 1 else {case['number']}
        hit = [r for r in candidates if r not in used and r.get('Case Number') in names]
        pairs += [(g, r) for r in hit]
        used += hit
        if not hit:
            unpaired.append(g)
    left = [r for r in candidates if r not in used]
    if len(unpaired) == 1 and left:
        pairs += [(unpaired[0], r) for r in left]
        unpaired, left = [], []
    pairs.sort(key=lambda p: (groups.index(p[0]), p[1].get('Case Number', '')))
    return pairs, unpaired, left


def _wanted(case, group, cells=None):
    """{field: value the workbook says RW should have} for one group; blank when the workbook has
    nothing to say (no result, unmapped type, unreadable or missing sections), so it isn't compared.
    Also 'Other Types' (the group's other reliefs' RW types, which also match) and 'Blank Order
    Status' (True when the result means RW's Order Status should be empty: Dismissed, Docketed).
    cells replaces the group's cells (whole sections for _WELL_TYPES)."""
    case_status, order_status = case['expected'] or ('', '')
    cells = group['cells'] if cells is None else cells
    sections_ok = cells and not group['unread']
    return {'Case Status': case_status, 'Order Status': order_status, 'Case Type': case['type'] or '',
            'Sections': _describe(cells) if sections_ok else '',
            'Hearing Date': _fmt_date(case['hearing']), 'Order Number': case['order_number'],
            'Date Order Signed': _fmt_date(case['signed']),
            'Other Types': [t for t in group['types'] if t != case['type']],
            'Blank Order Status': bool(case['expected']) and not order_status}


def _same(field, wanted, rec, well=False):
    """True when RW's value of field agrees with the workbook's (dates as dates, sections as land,
    by whole section when well, a type also by _TYPE_FAMILIES, the rest ignoring case and spacing)."""
    have = rec.get(field, '')
    if field == 'Sections':
        have_cells = _rw_cells(have)
        return (_whole_sections(have_cells) if well else have_cells) == _rw_cells(wanted)
    if field == 'Case Type' and have.lower() in {t.lower() for t in _TYPE_FAMILIES.get(wanted, ())}:
        return True
    if field in _DATE_FIELDS:
        return _date(have) == _date(wanted)
    return ' '.join(have.split()).lower() == ' '.join(wanted.split()).lower()


def _is_well(rec, case=None):
    """True when the RW record (or the workbook case) has one of the _WELL_TYPES."""
    return rec.get('Case Type', '') in _WELL_TYPES or bool(case and case['type'] in _WELL_TYPES)


def _differences(wanted, rec, fields, well=False):
    """Fields (of those the report has) where the workbook has a value and RW disagrees (well:
    compare sections whole, see _WELL_TYPES; a Case Type matching one of the Other Types counts as
    agreeing), plus Order Status when it should be blank and RW has one."""
    def differs(f):
        """True when RW's value of field f disagrees with the workbook's."""
        if not wanted[f]:
            return f == 'Order Status' and wanted['Blank Order Status'] and bool(rec.get(f, '').strip())
        values = [wanted[f]] + (wanted['Other Types'] if f == 'Case Type' else [])
        return not any(_same(f, v, rec, well) for v in values)
    return [f for f in fields if differs(f)]


def _pair_notes(want, rec, diffs):
    """Notes for one pairing: RW dates the script can't read (they always differ) and an Order
    Status RW should have blank (the upload never writes a blank, so it's cleared by hand)."""
    notes = [f'RW {f} "{rec[f]}" not read as a date' for f in diffs
             if f in _DATE_FIELDS and rec.get(f, '').strip() and _date(rec[f]) is None]
    if 'Order Status' in diffs and not want['Order Status']:
        notes.append(f'RW Order Status "{rec["Order Status"]}" should be blank ({want["Case Status"]}): '
                     'clear it in Ark by hand')
    return notes


def _rw_value(field, rec):
    """RW's own value of field for the upload file: a readable date as MM/DD/YYYY (so the file has
    one date format), anything else as the report gives it."""
    value = rec.get(field, '')
    got = _date(value) if field in _DATE_FIELDS else None
    return _fmt_date(got) if got else value


def _compare(cases, rw, fields):
    """(comparison rows, upload rows, Not in RW rows, Counter of Results, cases compared). In scope:
    cases with RW records and Phoenix cases; the rest are ignored. Each group gets a row per paired
    RW record (Match / Differs) or a Not in RW row. Not in workbook: leftover <number> / <number>-k sub-cases of a compared case (Roman-numeral field-rules sub-cases such as
    32665-V-1 skipped, since an order covers only a few) and every RW case whose number isn't in the
    workbook (Utah numbers skipped while UT isn't read)."""
    index = defaultdict(list)
    for rec in rw:
        index[rec.get('Case Number', '').split('-')[0]].append(rec)
    rows, uploads, missing, counts = [], {}, [], Counter()

    def add(fields_):
        """Append one comparison row from a dict of column values and count its Result."""
        rows.append([fields_.get(h, '') for h in _DIFF_HEADER])
        counts[fields_['Result']] += 1

    claimed, compared = set(), 0
    for case in cases:
        candidates = _candidates(case, index)
        claimed |= {id(r) for r in candidates}
        if not candidates and not case['phoenix']:
            continue
        compared += 1
        notes = list(case['notes'])
        if case['type_text'] and not case['type']:
            notes.append('Type not compared')
        base = {'Case Number': case['number'], 'Sheet': case['sheet'], 'State': case['state'],
                'County': case['county'], 'Workbook Operator': case['operator'], 'Notes': '; '.join(notes)}
        pairs, unpaired, left = _pair(case, candidates)
        many = len(case['groups']) > 1
        for k, g in enumerate(case['groups'], 1):
            wanted = _wanted(case, g, _whole_sections(g['cells']) if case['type'] in _WELL_TYPES else None)
            row = dict(base, Group=f'{k} of {len(case["groups"])}' if many else '',
                       Rows=', '.join(map(str, g['rows'])),
                       **{f'Workbook {f}': wanted[f] for f in _FIELDS})
            row['Workbook Case Type'] = case['type'] or case['type_text']
            row['Workbook Sections'] = wanted['Sections'] or g['text']
            for pg, rec in pairs:
                if pg is not g:
                    continue
                well = _is_well(rec, case)  # either side a well type: whole sections on both sides
                want = _wanted(case, g, _whole_sections(g['cells'])) if well else wanted
                diffs = _differences(want, rec, fields, well)
                mine = _rw_cells(want['Sections']) if 'Sections' in diffs else set()
                theirs = _rw_cells(rec.get('Sections')) if 'Sections' in diffs else set()
                if well:
                    theirs = _whole_sections(theirs)
                add(dict(row, **{'Workbook Sections': want['Sections'] or g['text']},
                         Result=_DIFFERS if diffs else _MATCH, Differences=', '.join(diffs),
                         **{'RW Case Number': rec.get('Case Number', ''), 'Record Id': rec['Record Id'],
                            'RW Applicant': rec.get('Applicant', '')},
                         **{f'RW {f}': rec.get(f, '') for f in _FIELDS},
                         **{'Missing from RW': _describe(mine - theirs), 'Extra in RW': _describe(theirs - mine)},
                         Notes='; '.join(filter(None, [row['Notes']] + _pair_notes(want, rec, diffs)))))
                fixes = [f for f in diffs if want[f]]  # an Order Status that should be blank is left to the user
                if fixes:
                    # Every column is filled: the workbook value where it differs, RW's own value
                    # elsewhere, so the Update never blanks a field Ark already has
                    upload = [rec['Record Id']] + [want[f] if f in fixes else _rw_value(f, rec) for f in fields]
                    if rec['Record Id'] in uploads and uploads[rec['Record Id']] != upload:
                        rows[-1][_DIFF_HEADER.index('Notes')] += '; Two workbook rows pair with this record: upload uses the first'
                    uploads.setdefault(rec['Record Id'], upload)
            if g in unpaired:
                add(dict(row, Result=_NOT_IN_RW))
                number = f'{case["number"]}-{k}' if many else case['number']
                missing.append([number, case['state'], case['operator']]
                               + [row[f'Workbook {f}'] if f in ('Case Type', 'Sections') else wanted[f]
                                  for f in _FIELDS])
        plain = re.compile(rf'{re.escape(case["number"])}(-\d+)?')
        for rec in left:
            if plain.fullmatch(rec.get('Case Number', '')):
                add(_rw_only(rec, 'No workbook row of this case pairs with it', case))
    for rec in sorted((r for r in rw if id(r) not in claimed), key=lambda r: r.get('Case Number', '')):
        if not _UTAH_NUMBER_RE.match(rec.get('Case Number', '')):
            add(_rw_only(rec, f'Case not on the {"/".join(_SHEETS)} sheets'))
    return rows, list(uploads.values()), missing, counts, compared


def _rw_only(rec, note, case=None):
    """A Not in workbook row for an RW record (case: the workbook case it belongs to, if any)."""
    fields = {'Result': _NOT_IN_WORKBOOK, 'RW Case Number': rec.get('Case Number', ''),
              'Record Id': rec['Record Id'], 'RW Applicant': rec.get('Applicant', ''), 'Notes': note,
              **{f'RW {f}': rec.get(f, '') for f in _FIELDS}}
    if case:
        fields.update({'Case Number': case['number'], 'Sheet': case['sheet'], 'State': case['state'],
                       'County': case['county'], 'Workbook Operator': case['operator']})
    return fields


def _summary(case_count, rw_count, counts, uploads, fields, paths):
    """Text of the finish message box: result counts, fields not compared and where the files went."""
    skipped = [f for f in _FIELDS if f not in fields]
    text = (f'Compared {case_count} workbook case(s) ({"/".join(_SHEETS)}) with {rw_count} RW case(s): '
            f'{counts[_MATCH]} match, {counts[_DIFFERS]} differ, {counts[_NOT_IN_RW]} case(s) / sub-case(s) not in RW, '
            f'{counts[_NOT_IN_WORKBOOK]} RW case(s) not in the workbook.\n\n'
            f'{ark_common.plural(len(uploads), "RW record")} to update. Nothing was uploaded: check the upload '
            f'file, then load it with the Ark Data Loader as a Cases Update.')
    if skipped:
        text += f'\n\nNot compared (no such column in the RW report): {", ".join(skipped)}'
    return text + '\n\nDifferences: {}\nUpload file: {}\nNot in RW: {}'.format(*paths)


def main(today=None):
    """Read the workbook, download the RW report, compare, write the three CSVs and show the summary.
    A missing or locked workbook, a missing report column or a CSV open in Excel gets its own box.
    today (for tests) decides which hearings are still to come."""
    today = today or date.today()
    out_dir = ark_common.read_folder('DocketUpdates', _DEFAULT_DIR)
    os.makedirs(out_dir, exist_ok=True)
    paths = [os.path.join(out_dir, name) for name in (_DIFF_FILE, _UPLOAD_FILE, _MISSING_FILE)]
    try:
        cases = _cases(_read_workbook(ark_common.read_setting('Files', 'HearingDockets', _DEFAULT_WORKBOOK)), today)
        report_path = os.path.join(out_dir, _REPORT_FILE)
        with ark_common.ark_page(_REPORT_URL) as page:
            ark_common.trigger_download(page, _REPORT_URL).save_as(report_path)
        rw, fields = _read_rw(report_path)
    except InputError as e:
        _notify(str(e), 0x10)
        return
    diff_rows, uploads, missing, counts, compared = _compare(cases, rw, fields)
    try:
        _write_csv(paths[0], _DIFF_HEADER, diff_rows)
        _write_csv(paths[1], ['Record Id'] + fields, uploads)
        _write_csv(paths[2], _MISSING_HEADER, missing)
    except PermissionError as e:
        _notify(f'Could not write {e.filename or paths[0]}.\n\nClose it in Excel and click Docket Updates again.',
                0x10)
        return
    _notify(_summary(compared, len(rw), counts, uploads, fields, paths))


if __name__ == '__main__':
    ark_common.run_button('Docket Updates', main,
                          os.path.join(ark_common.read_folder('DocketUpdates', _DEFAULT_DIR), 'docket_updates_error.log'))
