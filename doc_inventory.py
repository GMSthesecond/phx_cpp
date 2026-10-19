"""Scanned-document inventory shared by the "Index" button's three county scripts
(doc_inventory_roosevelt / _richland / _divide.py, which main.cpp runs in that order, and Mirror's
Index step). Not launched directly.

For one county: walks its image folders in Box Drive, skips PDFs already listed in column A of the
county's tab of the index workbook ([Files] IndexWorkbook), writes each new PDF's name, folder and
page count to <County>.csv in the [Folders] Index folder, then backs up the workbook as
"<workbook>.<County>.bak" and appends the same rows to the tab. Prints "@PROGRESS <done> <total>"
(total -1 when the county isn't pre-counted, for a marquee bar) and "@DONE <processed>" on stdout
for main.cpp's progress bar and summary. Problems a person can fix (workbook missing or open in
Excel, tab missing) stop with SystemExit(message); as a button the message is shown in an error
box (run_button). Requires openpyxl and PyPDF2.
"""
import csv
import os
import shutil
import sys
import time
from dataclasses import dataclass
from typing import Iterable, List, Optional, Set, Tuple

import ark_common

_BOX_IMAGES = r"C:\Cloud\Box\Land\Title Folder\Document Images"  # Box Drive root of every county's image folders
_DEFAULT_INDEX_DIR = r"C:\Users\Ethan Mesecher\Desktop\Index"  # CSV folder used when [Folders] Index is unset
_DEFAULT_WORKBOOK = r"C:\Cloud\Box\Ethan Mesecher\Index of Scanned Docs.xlsx"  # workbook used when [Files] IndexWorkbook is unset
EXCEL_COLUMN_INDEX = 1  # workbook column holding the indexed file names (A = 1)

# What to put in the Folder column:
#   "parent" -> the immediate parent directory of the file
#   "root"   -> the configured root label (e.g., "A Books")
FOLDER_NAME_MODE = "parent"

PDF_ONLY = True  # only PDFs are listed and page-counted


@dataclass
class County:
    """One county the Index button inventories."""
    name: str             # workbook tab and CSV name, e.g. "Roosevelt"
    roots: List[str]      # Box Drive folders scanned recursively
    match_stem: bool      # the tab lists names without ".pdf", so match on the stem (Richland)
    precount: bool        # count the new files first for a percentage bar (False: running tally, marquee)
    csv_bom: bool = False  # write the CSV as UTF-8 with BOM (Divide), so Excel opens it cleanly


def _images(*parts):
    """A folder under _BOX_IMAGES."""
    return os.path.join(_BOX_IMAGES, *parts)


# The three counties, in the order the Index button runs them
ROOSEVELT = County("Roosevelt", [_images("Montana", "Roosevelt County Images", d) for d in (
    "Books 1 through 250", "Books 251 through 499", "Books 500+ and Modern",
    "Misc, LP, Lien, Random Files, and Plats", "Older Docs with only doc numbers",
    "Transcribed from Sheridan County")], match_stem=False, precount=True)
RICHLAND = County("Richland", [_images("Montana", "Richland County Images", d) for d in (
    "A Books", "B Books", "C Books", "D Books", "E Books", "Doc Numbers Only", "Misc Docs",
    "TD Books", "WR Books")], match_stem=True, precount=False)
DIVIDE = County("Divide", [_images("North Dakota", "Divide County Images", d) for d in (
    "Deeds", "Misc", "Modern", "Mort", "Plat", "Probates and Judgements",
    "Unnamed and Speciality Books")], match_stem=False, precount=True, csv_bom=True)


def output_csv(county):
    """Where the county's CSV goes: <[Folders] Index>/<County>.csv."""
    return os.path.join(ark_common.read_folder('Index', _DEFAULT_INDEX_DIR), f"{county.name}.csv")


def workbook_path():
    """The index workbook: [Files] IndexWorkbook, else _DEFAULT_WORKBOOK."""
    return ark_common.read_setting('Files', 'IndexWorkbook', _DEFAULT_WORKBOOK)


def win_long(path: str) -> str:
    """Return an absolute path with the \\\\?\\ (or \\\\?\\UNC\\) prefix so Windows allows >260 chars."""
    if os.name == "nt":
        path = os.path.abspath(path)
        if not path.startswith("\\\\?\\"):
            if path.startswith("\\\\"):  # UNC path
                return "\\\\?\\UNC\\" + path[2:]
            return "\\\\?\\" + path
    return path


def _key(name: str, match_stem: bool) -> str:
    """The form a file name is compared in: trimmed, lower-cased, without extension if match_stem."""
    name = name.strip().lower()
    return os.path.splitext(name)[0].strip() if match_stem else name


def _stop_if_open_in_excel(excel_path: str) -> None:
    """Stop (SystemExit) if Excel has the workbook open: saving would fail after the long scan."""
    try:
        with open(excel_path, "r+b"):
            pass
    except PermissionError:
        sys.exit(f"Close {os.path.basename(excel_path)} in Excel, then run Index again.\n({excel_path})")


def load_indexed_names(excel_path: str, sheet: str, match_stem: bool, col_index: int = EXCEL_COLUMN_INDEX) -> Set[str]:
    """The names in the sheet's column (1-based), compared as _key does. Stops (SystemExit) if
    openpyxl, the workbook or the sheet is missing."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Index requires 'openpyxl'. Install with: py -m pip install openpyxl")

    if not os.path.exists(excel_path):
        sys.exit(f"Excel file not found: {excel_path}")

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        sys.exit(f"Sheet '{sheet}' not found in workbook. Found: {wb.sheetnames}")

    names: Set[str] = set()
    for (val,) in wb[sheet].iter_rows(min_col=col_index, max_col=col_index, values_only=True):
        if val is not None and str(val).strip():
            names.add(_key(str(val), match_stem))
    wb.close()
    return names


def iter_files(root: str) -> Iterable[Tuple[str, str]]:
    """Yield (dirpath, filename) for files under root, skipping folders that can't be read."""
    for dirpath, _dirnames, filenames in os.walk(root, onerror=lambda e: None):
        for fname in filenames:
            yield dirpath, fname


def _wanted(fname: str, indexed_names: Set[str], match_stem: bool) -> bool:
    """Whether a file should be listed: a PDF (when PDF_ONLY) not already in the workbook."""
    if PDF_ONLY and not fname.lower().endswith(".pdf"):
        return False
    return _key(fname, match_stem) not in indexed_names


def count_files_to_process(roots: List[str], indexed_names: Set[str], match_stem: bool) -> int:
    """How many files under roots _wanted would list (for the percentage bar)."""
    return sum(1 for root in roots for _, fname in iter_files(root) if _wanted(fname, indexed_names, match_stem))


class Progress:
    """Throttled progress reporter for both a human reader and the C++ launcher.

    Each report prints a human line and a machine-readable "@PROGRESS <done> <total>" line that
    main.cpp (RunOneIndexScript) parses; total is -1 when not pre-counted (marquee bar)."""

    def __init__(self, total: Optional[int], min_interval: float = 0.5):
        """total is the pre-counted number of files, or None for a running tally."""
        self.total = total
        self.done = 0
        self.last_t = 0.0
        self.min_interval = min_interval

    def tick(self, n: int = 1):
        """Add n to the done count and report at most every min_interval seconds."""
        self.done += n
        now = time.time()
        if now - self.last_t >= self.min_interval:
            self.last_t = now
            if self.total is None:
                print(f"\rProcessed so far: {self.done:,}", end="", flush=True)
                print(f"@PROGRESS {self.done} -1", flush=True)
            else:
                pct = (self.done / self.total * 100) if self.total > 0 else 0.0
                print(f"\rProcessing: {self.done:,}/{self.total:,} ({pct:5.1f}%)", end="", flush=True)
                print(f"@PROGRESS {self.done} {self.total}", flush=True)

    def finish(self):
        """Print the final human-readable line (no @PROGRESS tag)."""
        print(f"\rProcessed: {self.done:,}")


def get_pdf_page_count(pdf_path: str) -> Optional[str]:
    """The PDF's page count as a string, a status like "Encrypted/Unreadable", or None if the file vanished."""
    try:
        from PyPDF2 import PdfReader
        with open(pdf_path, "rb") as f:
            reader = PdfReader(f)
            try:
                _ = len(reader.pages)
            except Exception:
                try:
                    reader.decrypt("")  # some encrypted PDFs open with an empty password
                except Exception:
                    return "Encrypted/Unreadable"
            try:
                return str(len(reader.pages))
            except Exception:
                return "Encrypted/Unreadable"
    except FileNotFoundError:
        return None  # file disappeared mid-scan; leave blank
    except PermissionError:
        return "No Permission"
    except Exception:
        return "Unreadable"


def folder_label_for_path(full_path: str, root_label: str) -> str:
    """The Folder column: the file's parent folder name ("parent" mode) or the root's label ("root")."""
    if FOLDER_NAME_MODE == "root":
        return root_label
    parent = os.path.basename(os.path.dirname(full_path.rstrip("\\/")))
    return parent or root_label


def backup_path(excel_path: str, county: County) -> str:
    """The county's own workbook backup, so one county's backup never overwrites another's."""
    return f"{excel_path}.{county.name}.bak"


def append_rows_to_excel(excel_path: str, county: County, rows: List[List[str]]) -> None:
    """Back up the workbook (backup_path), then append rows to the county's tab. Stops
    (SystemExit) with a readable message if Excel has the workbook open."""
    if not rows:
        print("No new rows to add to the Excel index.")
        return

    from openpyxl import load_workbook

    shutil.copy2(excel_path, backup_path(excel_path, county))

    wb = load_workbook(excel_path)
    if county.name not in wb.sheetnames:
        print(f"Sheet '{county.name}' not found in workbook; skipping Excel update.")
        wb.close()
        return

    ws = wb[county.name]
    for fname, folder_val, page_count in rows:
        # Keep PageCount numeric (matches existing rows) unless it's a status like "Encrypted/Unreadable"
        ws.append([fname, folder_val, int(page_count) if str(page_count).isdigit() else page_count])
    try:
        wb.save(excel_path)
    except PermissionError:
        sys.exit(f"Could not save {os.path.basename(excel_path)}: it was opened in Excel during the run.\n"
                 f"The new {county.name} rows are in {output_csv(county)}. Close Excel and run Index again.")
    finally:
        wb.close()
    print(f"Added {len(rows):,} row(s) to the '{county.name}' tab in {excel_path}")


def run(county: County) -> int:
    """Inventory one county: scan its roots for unindexed PDFs, write them to output_csv, then
    append them to the workbook tab. Prints "@DONE <processed>" for main.cpp's summary and
    returns that count (Mirror's Index step adds the counties up)."""
    excel_path = workbook_path()
    out_csv = output_csv(county)
    indexed_names = load_indexed_names(excel_path, county.name, county.match_stem)
    _stop_if_open_in_excel(excel_path)  # before the long scan, not after it

    prepared_roots = [(win_long(r), os.path.basename(r.rstrip("\\/")) or r) for r in county.roots]
    total = (count_files_to_process([lr for lr, _ in prepared_roots], indexed_names, county.match_stem)
             if county.precount else None)
    prog = Progress(total)

    processed = 0
    new_rows: List[List[str]] = []
    os.makedirs(os.path.dirname(out_csv), exist_ok=True)
    with open(out_csv, mode="w", newline="", encoding="utf-8-sig" if county.csv_bom else "utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["FileName", "Folder", "PageCount"])
        for long_root, root_label in prepared_roots:
            for dirpath, fname in iter_files(long_root):
                if not _wanted(fname, indexed_names, county.match_stem):
                    continue
                full_path = os.path.join(dirpath, fname)
                try:
                    page_count = (get_pdf_page_count(full_path) or "") if fname.lower().endswith(".pdf") else ""
                    row = [fname, folder_label_for_path(full_path, root_label), page_count]
                    writer.writerow(row)
                    new_rows.append(row)
                    processed += 1
                except Exception:
                    pass  # one bad file never stops the county
                finally:
                    prog.tick(1)

    prog.finish()
    print(f"@DONE {processed}", flush=True)
    print(f"Report written to: {out_csv}")

    append_rows_to_excel(excel_path, county, new_rows)
    return processed


def run_button(county: County) -> int:
    """A county script launched by the Index button: run it, and show any stop or error in a box
    titled "Index (<County>)" (log: index_<county>_error.log in the Index folder)."""
    def checked():
        try:
            return run(county)
        except SystemExit as e:
            if isinstance(e.code, str):
                raise RuntimeError(e.code) from None  # a readable stop: show the message itself
            raise
    log = os.path.join(ark_common.read_folder('Index', _DEFAULT_INDEX_DIR), f"index_{county.name.lower()}_error.log")
    return ark_common.run_button(f"Index ({county.name})", checked, log)
