"""TPP vs. LHPP price comparison, launched by the "TPP-LHPP" button.

Downloads the Offers and LandHoldingTransactions reports from ark.phoenixenergy.com into
the TPP-LHPP folder, sums LHT values per offer, and writes offers whose Total Purchase
Price differs from their LHT total by more than $1 to PHX_Price_Paid_Export.xlsx, then
shows a message box saying which exported offers (if any) have a Close Date.
Also Mirror's "TPP-LHPP" step (mirror_step): no message box, and the report line counts only
mismatches closing in the next two weeks.
Uses ark_common (credentials, login, session.json), Playwright and openpyxl; errors are
shown in a message box since the launcher gives the script no console.
"""
import csv
import os
from datetime import date, timedelta
import ark_common
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.cell import WriteOnlyCell


def _notify(message):
    """Show an information message box titled TPP-LHPP (errors and the Close Date summary)."""
    ark_common.notify('TPP-LHPP', message, 0x40)


# Download/output folder ([Folders] TPPLHPP, set in Settings)
_OUT_DIR = ark_common.read_folder('TPPLHPP', r"C:\Users\Ethan Mesecher\Desktop\TPP-LHPP")

# (url, generation_wait_ms) — the LandHoldingTransactions report has ~350k+ rows and can take
# well over the 10s default to finish generating server-side; grabbing the download too early
# silently returns a stale previously-generated copy instead of erroring, which was causing
# offers with recently-added LandHoldingTransactions to show up as having 0 LHT total.
_REPORT_URLS = [
    ('https://ark.phoenixenergy.com/report?recordId=688d3dfda95ef3a2a6ba14f9', 15_000),   # Offers report
    ('https://ark.phoenixenergy.com/report?recordId=689257e3d2b183ad4f221902', 120_000),  # LandHoldingTransactions — large, slow to generate
]

_OFFERS_FILENAME = 'Offers_Report_For_Price_Comparison_EM.csv'  # expected name of the downloaded Offers report
_LHT_FILENAME = 'LandHoldingTransactions_For_Price_Comparison_EM.csv'  # expected name of the downloaded LHT report
_EXPORT_FILENAME = 'PHX_Price_Paid_Export.xlsx'  # mismatch workbook written to _OUT_DIR
_MAX_LISTED = 25  # most offers named in the Close Date message box; the rest are only counted
_SOON_DAYS = 14  # Mirror counts mismatches closing from today through this many days out

# Highlight applied to every cell of an exported mismatch row
_YELLOW_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')


def _download_files():
    """Log in to Ark and download both _REPORT_URLS reports into _OUT_DIR.

    Waits 30s between downloads; always saves session.json and closes the browser.
    """
    os.makedirs(_OUT_DIR, exist_ok=True)

    with ark_common.ark_page(_REPORT_URLS[0][0]) as page:
        for i, (url, generation_wait_ms) in enumerate(_REPORT_URLS):
            if i > 0:
                page.wait_for_timeout(30_000)  # 30-second delay between downloads

            dl = ark_common.trigger_download(page, url, generation_wait_ms=generation_wait_ms)
            dest = os.path.join(_OUT_DIR, dl.suggested_filename)
            dl.save_as(dest)
            print(f'Saved: {dest}', flush=True)

        page.wait_for_timeout(10_000)  # let the last download finish writing before closing


def _find_file(filename):
    """Return the path of filename in _OUT_DIR (case-insensitive match), or None."""
    for name in os.listdir(_OUT_DIR):
        if name.lower() == filename.lower():
            return os.path.join(_OUT_DIR, name)
    return None


def _process_offers(offers):
    """Fill offers (dict keyed by offer Id) with TPP, status and close date from the Offers CSV.

    Silently does nothing if the Offers file isn't found.
    """
    path = _find_file(_OFFERS_FILENAME)
    if not path:
        return

    with open(path, newline='', encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            offer_id = row['Id'].strip()
            offer = offers.setdefault(offer_id, {
                'offer_tpp': 0.0, 'lht_tpp': 0.0, 'status': '', 'close_date': '',
            })

            tpp = row['Total Purchase Price'].strip()
            if tpp:
                try:
                    offer['offer_tpp'] = float(tpp)
                except ValueError:
                    print(f'Error parsing {offer_id}')

            status = row['Status'].strip()
            if status:
                offer['status'] = status

            close_date = row['Close Date'].strip()
            if close_date:
                offer['close_date'] = close_date


def _process_lht(offers):
    """Sum LHT values per offer Id and store them as lht_tpp on offers already in the dict.

    Silently does nothing if the LHT file isn't found.
    """
    path = _find_file(_LHT_FILENAME)
    if not path:
        return

    # Multiple LandHoldingTransaction rows can share the same offer id — sum their values.
    totals = {}
    with open(path, newline='', encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            offer_id = row['LandHoldingTransaction: Offer Id'].strip()
            value = row['LandHoldingTransaction: Total LandHolding Value'].strip()
            if not value:
                continue
            try:
                totals[offer_id] = totals.get(offer_id, 0.0) + float(value)
            except ValueError:
                print(f'Error parsing {offer_id}')

    # Only apply LHT totals to offers that actually appear in the Offers report —
    # LandHoldingTransactions with no matching offer aren't a real comparison and
    # would otherwise flood the export as trivial "mismatches" (0 vs. some value).
    for offer_id, lht_tpp in totals.items():
        if offer_id in offers:
            offers[offer_id]['lht_tpp'] = lht_tpp


def _export(offers):
    """Write offers whose LHT total and TPP differ by more than $1 to the yellow-highlighted export workbook.

    Returns the exported offer Ids in row order.
    """
    export_path = os.path.join(_OUT_DIR, _EXPORT_FILENAME)

    # write_only mode streams rows instead of building an in-memory cell grid —
    # required here since a mismatched offer highlights all 5 cells in its row,
    # and normal openpyxl cell styling is far too slow at 100k+ styled rows.
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Export Data')

    ws.append(['Offer ID', 'Total Purchase Price', 'LHT Price Total', 'Close Date', 'Status'])

    exported = []
    for offer_id, offer in offers.items():
        diff = offer['lht_tpp'] - offer['offer_tpp']
        # A >$1 tolerance filters out float-rounding noise from summing many small
        # LHT values (e.g. off by a few cents) that isn't a real price mismatch.
        if -1 <= diff <= 1:
            continue

        values = [offer_id, offer['offer_tpp'], offer['lht_tpp'], offer['close_date'], offer['status']]
        row = []
        for value in values:
            cell = WriteOnlyCell(ws, value=value)
            cell.fill = _YELLOW_FILL
            row.append(cell)
        ws.append(row)
        exported.append(offer_id)

    wb.save(export_path)
    print(f'Data successfully exported to: {export_path}', flush=True)
    return exported


def _close_date_message(offers, exported):
    """Message box text saying how many exported offers have a Close Date, listing up to _MAX_LISTED."""
    if not exported:
        return 'No offers were exported (every price matches within $1).'

    plural = 's' if len(exported) != 1 else ''
    dated = [offer_id for offer_id in exported if offers[offer_id]['close_date']]
    if not dated:
        return f'No Close Date on the {len(exported)} exported offer{plural}.'

    lines = [f'{offer_id}  {offers[offer_id]["close_date"]}' for offer_id in dated[:_MAX_LISTED]]
    if len(dated) > _MAX_LISTED:
        lines.append(f'...and {len(dated) - _MAX_LISTED} more (see {_EXPORT_FILENAME}).')
    return f'Close Date found on {len(dated)} of {len(exported)} exported offer{plural}:\n\n' + '\n'.join(lines)


def _closing_soon(offers, exported):
    """(soon, unreadable): how many exported offers have a Close Date from today through _SOON_DAYS
    from now, and how many have a Close Date ark_common.parse_date can't read (blank dates are neither)."""
    today = date.today()
    soon = unreadable = 0
    for offer_id in exported:
        text = offers[offer_id]['close_date'].strip()
        if not text:
            continue
        try:
            close_date = ark_common.parse_date(text)
        except ValueError:
            unreadable += 1
            continue
        if today <= close_date <= today + timedelta(days=_SOON_DAYS):
            soon += 1
    return soon, unreadable


def main(quiet=False):
    """Download both reports, merge offers with their LHT totals, export the mismatches,
    and show which exported offers have a Close Date (no message box when quiet, for Mirror).
    Returns (offers, exported offer Ids)."""
    _download_files()

    offers = {}
    _process_offers(offers)
    _process_lht(offers)
    exported = _export(offers)
    if not quiet:
        _notify(_close_date_message(offers, exported))
    return offers, exported


def mirror_step():
    """Mirror's TPP-LHPP step: run quietly and return (result kind, report line) counting
    mismatches closing in the next two weeks, plus any whose Close Date couldn't be read (they
    might be closing soon too)."""
    soon, unreadable = _closing_soon(*main(quiet=True))
    review = f': review {_EXPORT_FILENAME}'
    odd = "Close Date Mirror couldn't read"
    if soon:
        line = f'{ark_common.plural(soon, "price mismatch", "price mismatches")} closing in the next two weeks'
        if unreadable:
            line += f', and {unreadable} with a {odd}'
        return ark_common.REVIEW, line + review
    if unreadable:
        have = 'has' if unreadable == 1 else 'have'
        return ark_common.REVIEW, (f'No price mismatches closing in the next two weeks, but {unreadable} '
                                   f'{have} a {odd}{review}')
    return ark_common.NOTHING, 'No price mismatches closing in the next two weeks'


if __name__ == '__main__':
    ark_common.run_button('TPP-LHPP', main, os.path.join(_OUT_DIR, 'tpp_lhpp_error.log'))
